from __future__ import annotations

import datetime as dt
import json
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from analystos_engine.ingest import (
    IngestError,
    IngestOptions,
    infer_column,
    ingest_file,
    inspect_file,
    normalize_column_name,
)
from analystos_engine.store import WorkspaceStore


@pytest.fixture()
def store():
    s = WorkspaceStore.in_memory()
    yield s
    s.close()


def _inject_cached_values(path: Path, values: dict[str, str]) -> None:
    """openpyxl cannot write cached formula results; patch the sheet XML like Excel would."""
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode()
                for ref, val in values.items():
                    xml = xml.replace(f'<c r="{ref}"><f>', f'<c r="{ref}" t="n"><f>', 1)
                    start = xml.index(f'<c r="{ref}"')
                    vpos = xml.index("<v />", start)
                    xml = xml[:vpos] + f"<v>{val}</v>" + xml[vpos + 5 :]
                data = xml.encode()
            zout.writestr(item, data)
    tmp.replace(path)


def messy_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget 2026"
    # merged title row + subtitle + blank row, then header on row 4
    ws["A1"] = "Summit Supply Co. - Branch Budget"
    ws.merge_cells("A1:E1")
    ws["A2"] = "Prepared by FP&A"
    ws.append([])  # row 3 blank
    ws.append(["Branch", "Month", "Budget ($)", "Growth %", "Notes"])  # row 4
    rows = [
        ["Dallas", dt.datetime(2026, 7, 1), 125000, "3.5%", "ok"],
        ["Houston", dt.datetime(2026, 7, 1), 98000.5, "2%", None],
        [None, None, None, None, None],  # blank row inside the table
        ["Austin", dt.datetime(2026, 8, 1), 76000, "1.25%", "new store"],
    ]
    for r in rows:
        ws.append(r)
    ws.append(["Total", None, None, None, None])  # row 9
    ws["C9"] = "=SUM(C5:C8)"
    # second table on the same sheet, separated by a blank column (G onwards)
    ws["G4"] = "Region"
    ws["H4"] = "Manager"
    ws["G5"] = "North"
    ws["H5"] = "Kim"
    ws["G6"] = "South"
    ws["H6"] = "Lee"

    ws2 = wb.create_sheet("Serials")
    ws2.append(["order_date", "amount", "code", "mixed"])
    ws2.append([46235, 10, "00123", 5])
    ws2.append([46236, 20, "00456", "n/a"])
    ws2.append([46237, 30, "00789", "seven"])
    ws2.append([46238, "=B2+B3", "01000", 8])

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["a", "b"])
    hidden.append([1, 2])

    empty = wb.create_sheet("Empty")
    empty["A1"] = None
    wb.save(path)
    _inject_cached_values(path, {"C9": "299000.5"})
    return path


def test_excel_inspection(tmp_path):
    insp = inspect_file(messy_workbook(tmp_path / "budget.xlsx"))
    assert insp.format == "excel"
    names = [s.name for s in insp.sheets]
    assert names == ["Budget 2026", "Serials", "Hidden", "Empty"]
    assert insp.sheets[2].visible is False
    assert insp.sheets[3].empty
    assert any("hidden" in w for w in insp.warnings)

    budget_tables = [t for t in insp.tables if t.sheet == "Budget 2026"]
    assert len(budget_tables) == 2  # two table blocks on one sheet
    main = budget_tables[0]
    assert main.header_row == 4
    assert main.first_data_row == 5
    assert [c.name for c in main.columns] == ["branch", "month", "budget", "growth_pct", "notes"]
    types = {c.name: c.inferred_type for c in main.columns}
    assert types == {
        "branch": "string",
        "month": "date",
        "budget": "float",
        "growth_pct": "float",
        "notes": "string",
    }
    assert main.row_count == 3  # blank row removed, total row excluded
    reasons = {s.reason for s in main.skipped_rows}
    assert "total" in reasons
    assert main.title and "Summit Supply" in main.title
    assert main.preview_rows[0][0] == "Dallas"
    assert main.preview_rows[0][3] == pytest.approx(0.035)
    second = budget_tables[1]
    assert [c.name for c in second.columns] == ["region", "manager"]
    assert second.row_count == 2


def test_excel_serials_formulas_mixed(tmp_path):
    insp = inspect_file(messy_workbook(tmp_path / "budget.xlsx"))
    t = insp.table("Serials")
    types = {c.name: c for c in t.columns}
    assert types["order_date"].inferred_type == "date"
    assert any("serial" in n for n in types["order_date"].notes)
    assert t.preview_rows[0][0] == dt.date(1899, 12, 30) + dt.timedelta(days=46235)
    assert types["code"].inferred_type == "string"  # leading zeros preserved
    assert t.preview_rows[0][2] == "00123"
    mixed = types["mixed"]
    assert mixed.inferred_type == "string"
    assert mixed.candidate_type == "float"
    assert "seven" in mixed.nonconforming_examples
    assert mixed.null_count == 1  # "n/a"
    # formula without a cached value is reported
    assert any("formula cells without cached values" in w for w in insp.warnings)


def test_excel_formula_cached_value_read(tmp_path, store):
    path = messy_workbook(tmp_path / "budget.xlsx")
    # the Total row holds a formula with a cached value; including it via header override shows it
    wb = openpyxl.load_workbook(path, data_only=True)
    assert wb["Budget 2026"]["C9"].value == 299000.5


def test_excel_ingest_selected_table(tmp_path, store):
    path = messy_workbook(tmp_path / "budget.xlsx")
    insp = inspect_file(path)
    key = insp.tables[1].key
    info = ingest_file(store, path, IngestOptions(table_key=key, table_name="managers"))
    assert info.row_count == 2
    info = ingest_file(store, path, IngestOptions(sheet="Budget 2026", table_name="budget"))
    assert info.row_count == 3
    types = {c.name: c.type for c in info.columns}
    assert types["month"] == "DATE" and types["budget"] == "DOUBLE"
    assert store.execute_read("SELECT sum(budget) FROM budget").scalar() == pytest.approx(299000.5)
    info = ingest_file(
        store, path, IngestOptions(sheet="Serials", table_name="serials", column_types={"mixed": "float"})
    )
    vals = store.execute_read("SELECT mixed FROM serials ORDER BY order_date").column_values("mixed")
    assert vals == [5.0, None, None, 8.0]
    with pytest.raises(IngestError):
        ingest_file(store, path, IngestOptions(sheet="Nope"))
    with pytest.raises(IngestError):
        ingest_file(store, path, IngestOptions(table_key="Budget 2026!Z1:Z2"))


def test_excel_header_override(tmp_path, store):
    path = messy_workbook(tmp_path / "budget.xlsx")
    info = ingest_file(store, path, IngestOptions(sheet="Budget 2026", header_row=4, table_name="b2"))
    assert info.row_count == 3


def test_excel_source_not_modified(tmp_path, store):
    path = messy_workbook(tmp_path / "budget.xlsx")
    before = path.read_bytes()
    inspect_file(path)
    ingest_file(store, path, IngestOptions(sheet="Serials", table_name="s"))
    assert path.read_bytes() == before


def test_excel_no_header_numeric_block(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(5):
        ws.append([r, r * 2.5, r * 3])
    p = tmp_path / "nums.xlsx"
    wb.save(p)
    t = inspect_file(p).tables[0]
    assert t.header_row is None
    assert [c.name for c in t.columns] == ["column_1", "column_2", "column_3"]
    assert t.row_count == 5


def test_not_a_workbook(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(IngestError):
        inspect_file(p)


def test_legacy_xls_rejected(tmp_path):
    p = tmp_path / "old.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(IngestError, match="xlsx"):
        inspect_file(p)


# ----------------------------------------------------------------------------- CSV


def test_csv_with_title_rows_semicolon_and_latin1(tmp_path, store):
    p = tmp_path / "Sales Export.csv"
    content = (
        "Sales export generated 2026-09-01\n"
        "\n"
        "Order ID;Customer;Amount;Order Date;Active\n"
        "1;Café Olé;1.234,50;2026-08-01;yes\n"
    )
    # European decimal comma is ambiguous; use plain values below
    content = (
        "Sales export generated 2026-09-01\n"
        "\n"
        "Order ID;Customer;Amount;Order Date;Active\n"
        '1;Café Olé;"$1,234.50";08/01/2026;yes\n'
        "2;Zoë Ltd;99;08/15/2026;no\n"
        "3;Plain;;08/20/2026;yes\n"
    )
    p.write_bytes(content.encode("latin-1"))
    insp = inspect_file(p)
    assert insp.dialect is not None and insp.dialect.delimiter == ";"
    assert insp.encoding and insp.encoding.lower().replace("_", "-") not in ("utf-8",)
    t = insp.tables[0]
    assert t.header_row == 3
    types = {c.name: c.inferred_type for c in t.columns}
    assert types == {
        "order_id": "integer",
        "customer": "string",
        "amount": "float",
        "order_date": "date",
        "active": "boolean",
    }
    info = ingest_file(store, p, IngestOptions())
    assert info.name == "sales_export"
    rec = store.execute_read("SELECT * FROM sales_export ORDER BY order_id").to_records()
    assert rec[0]["customer"] == "Café Olé"
    assert rec[0]["amount"] == 1234.5
    assert rec[2]["amount"] is None
    assert rec[1]["order_date"] == dt.date(2026, 8, 15)


def test_csv_utf8_bom_tab(tmp_path, store):
    p = tmp_path / "data.tsv"
    p.write_bytes("﻿id\tname\n1\ta\n2\tb\n".encode())
    insp = inspect_file(p)
    assert insp.dialect.delimiter == "\t" and insp.dialect.has_bom
    assert [c.name for c in insp.tables[0].columns] == ["id", "name"]
    ingest_file(store, p, IngestOptions(table_name="tsv"))
    assert store.row_count("tsv") == 2


def test_csv_malformed_dates_kept_as_text(tmp_path):
    p = tmp_path / "orders.csv"
    p.write_text("id,order_date\n1,2026-08-01\n2,2026-08-02\n3,2026-13-45\n4,not a date\n")
    col = inspect_file(p).tables[0].columns[1]
    assert col.inferred_type == "string"
    assert col.candidate_type == "date"
    assert set(col.nonconforming_examples) == {"2026-13-45", "not a date"}


def test_csv_ragged_rows_and_duplicates(tmp_path):
    p = tmp_path / "rag.csv"
    p.write_text("a,b,b\n1,2,3\n4,5\n6,7,8,9\n")
    t = inspect_file(p).tables[0]
    assert [c.name for c in t.columns][:3] == ["a", "b", "b_2"]
    assert any("different number of fields" in n for n in t.notes)


def test_csv_percent_accounting_negative(tmp_path):
    p = tmp_path / "fin.csv"
    p.write_text('line,delta,rate\nA,"(1,200)",5%\nB,300,10%\n')
    t = inspect_file(p).tables[0]
    assert t.preview_rows[0][1] == -1200.0
    assert t.preview_rows[1][2] == pytest.approx(0.10)


def test_csv_whitespace_and_na_tokens(tmp_path):
    p = tmp_path / "ws.csv"
    p.write_text("id,val\n  7 ,N/A\n8,  12 \n9,#DIV/0!\n")
    t = inspect_file(p).tables[0]
    assert t.columns[0].inferred_type == "integer"
    assert t.columns[1].inferred_type == "integer"
    assert t.columns[1].null_count == 2


def test_empty_and_missing_files(tmp_path):
    p = tmp_path / "empty.csv"
    p.write_text("")
    with pytest.raises(IngestError):
        inspect_file(p)
    with pytest.raises(IngestError):
        inspect_file(tmp_path / "missing.csv")
    with pytest.raises(IngestError):
        inspect_file(tmp_path)
    q = tmp_path / "x.exe"
    q.write_text("x")
    with pytest.raises(IngestError):
        inspect_file(q)


def test_size_limit(tmp_path):
    p = tmp_path / "big.csv"
    p.write_text("a\n" + "1\n" * 1000)
    with pytest.raises(IngestError, match="limit"):
        inspect_file(p, max_bytes=100)


def test_append_and_fail_modes(tmp_path, store):
    p = tmp_path / "t.csv"
    p.write_text("a\n1\n2\n")
    ingest_file(store, p, IngestOptions(table_name="t"))
    with pytest.raises(IngestError):
        ingest_file(store, p, IngestOptions(table_name="t"))
    ingest_file(store, p, IngestOptions(table_name="t", if_exists="append"))
    assert store.row_count("t") == 4
    with pytest.raises(IngestError):
        ingest_file(store, p, IngestOptions(table_name="bad name"))


def test_column_renames(tmp_path, store):
    p = tmp_path / "r.csv"
    p.write_text("a,b\n1,2\n")
    info = ingest_file(store, p, IngestOptions(table_name="r", column_names={"a": "Alpha"}))
    assert [c.name for c in info.columns] == ["alpha", "b"]
    with pytest.raises(IngestError):
        ingest_file(store, p, IngestOptions(table_name="r2", column_names={"zzz": "x"}))


# ---------------------------------------------------------------------- JSON / Parquet


def test_json_records_wrapped_and_nested(tmp_path, store):
    p = tmp_path / "customers.json"
    p.write_text(
        json.dumps(
            {
                "meta": {"n": 2},
                "data": [
                    {
                        "id": 1,
                        "name": "A",
                        "address": {"city": "Dallas"},
                        "tags": ["x"],
                        "since": "2024-01-05",
                    },
                    {"id": 2, "name": "B", "address": {"city": "Austin"}, "tags": [], "since": "2025-02-10"},
                ],
            }
        )
    )
    t = inspect_file(p).tables[0]
    names = {c.name: c.inferred_type for c in t.columns}
    assert names["address_city"] == "string"
    assert names["since"] == "date"
    assert names["tags"] == "string"
    info = ingest_file(store, p, IngestOptions())
    assert info.name == "customers" and info.row_count == 2


def test_ndjson_and_columnar(tmp_path, store):
    p = tmp_path / "events.jsonl"
    p.write_text('{"a": 1, "b": "x"}\n{"a": 2, "b": "y"}\n')
    assert inspect_file(p).format == "ndjson"
    q = tmp_path / "events2.json"
    q.write_text('{"a": 1}\n{"a": 2}\n')
    assert inspect_file(q).format == "ndjson"
    c = tmp_path / "cols.json"
    c.write_text(json.dumps({"x": [1, 2, 3], "y": ["a", "b", "c"]}))
    assert inspect_file(c).tables[0].row_count == 3
    bad = tmp_path / "bad.json"
    bad.write_text("[1, {")
    with pytest.raises(IngestError):
        inspect_file(bad)
    badl = tmp_path / "bad.jsonl"
    badl.write_text('{"a":1}\nnope\n')
    with pytest.raises(IngestError):
        inspect_file(badl)


def test_parquet(tmp_path, store):
    p = tmp_path / "Lines.parquet"
    pq.write_table(
        pa.table({"Line ID": [1, 2, 3], "amount": [1.5, 2.5, None], "d": [dt.date(2026, 1, 1)] * 3}), p
    )
    insp = inspect_file(p)
    t = insp.tables[0]
    assert t.row_count == 3
    assert [c.name for c in t.columns] == ["line_id", "amount", "d"]
    assert [c.inferred_type for c in t.columns] == ["integer", "float", "date"]
    info = ingest_file(store, p, IngestOptions())
    assert info.name == "lines"
    assert store.execute_read("SELECT sum(amount) FROM lines").scalar() == 4.0


# ------------------------------------------------------------------------ inference unit


def test_normalize_column_name():
    assert normalize_column_name("Revenue ($)", 0) == "revenue"
    assert normalize_column_name("Order Date", 0) == "order_date"
    assert normalize_column_name("", 2) == "column_3"
    assert normalize_column_name("2026 Budget", 0) == "c_2026_budget"
    assert normalize_column_name("Margin %", 0) == "margin_pct"


@pytest.mark.parametrize(
    ("values", "typ"),
    [
        ([1, 2, 3], "integer"),
        ([1.5, 2, None], "float"),
        (["true", "False", "yes"], "boolean"),
        (["1", "2", "x"], "string"),
        ([dt.datetime(2026, 1, 1, 10, 30), dt.datetime(2026, 1, 2)], "timestamp"),
        ([dt.date(2026, 1, 1), "2026-01-02"], "date"),
        (["2026-01-01 10:00:00", "2026-01-02 11:00:00"], "timestamp"),
        ([None, None], "string"),
        ([True, False], "boolean"),
        (["20260101", "20260102"], "integer"),
        (["1e3", "2.5"], "float"),
    ],
)
def test_infer_column(values, typ):
    _, info = infer_column(values, "col")
    assert info.inferred_type == typ


def test_yyyymmdd_in_date_named_column():
    _, info = infer_column(["20260101", "20260102"], "order_date")
    assert info.inferred_type == "date"


def test_infer_force_string_and_integer():
    s, info = infer_column(["1", "2"], "x", force="string")
    assert info.inferred_type == "string" and list(s) == ["1", "2"]
    s, info = infer_column(["1", "b"], "x", force="integer")
    assert info.inferred_type == "integer" and info.nonconforming_count == 1


def test_frame_roundtrip_types(tmp_path, store):
    p = tmp_path / "types.csv"
    df = pd.DataFrame({"i": [1, 2], "f": [1.5, None], "d": ["2026-01-01", "2026-01-02"], "b": ["yes", "no"]})
    df.to_csv(p, index=False)
    info = ingest_file(store, p, IngestOptions(table_name="types"))
    assert {c.name: c.type for c in info.columns} == {
        "i": "BIGINT",
        "f": "DOUBLE",
        "d": "DATE",
        "b": "BOOLEAN",
    }


def test_header_override_limits_width(tmp_path):
    from analystos_engine.ingest import load_table_frame

    path = messy_workbook(tmp_path / "budget.xlsx")
    df, table, _ = load_table_frame(path, IngestOptions(sheet="Budget 2026", header_row=4))
    assert list(df.columns) == ["branch", "month", "budget", "growth_pct", "notes"]
    assert table.header_row == 4


def test_header_override_beyond_sheet(tmp_path):
    path = messy_workbook(tmp_path / "budget.xlsx")
    with pytest.raises(IngestError):
        inspect_file(path)  # sanity: inspection works
        ingest_file(WorkspaceStore.in_memory(), path, IngestOptions(sheet="Serials", header_row=99))


def test_year_header_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Targets"
    ws.append(["Branch", 2025, 2026])
    ws.append(["Dallas", 100, 120])
    ws.append(["Austin", 80, 90])
    p = tmp_path / "targets.xlsx"
    wb.save(p)
    t = inspect_file(p).tables[0]
    assert t.header_row == 1
    assert [c.name for c in t.columns] == ["branch", "c_2025", "c_2026"]
    assert t.name_suggestion == "targets"


# ------------------------------------------------------------------ review R-08 / R-35


def test_inspection_honours_header_row_and_sheet_like_ingest(tmp_path, store):
    path = messy_workbook(tmp_path / "budget.xlsx")
    opts = IngestOptions(sheet="Budget 2026", header_row=2, table_name="b3")
    insp = inspect_file(path, options=opts, preview_rows=2)
    assert [s.name for s in insp.sheets] == ["Budget 2026"]
    table = insp.tables[0]
    assert table.header_row == 2
    assert len(table.preview_rows) <= 2
    info = ingest_file(store, path, opts)
    assert [c.name for c in table.columns] == [c.name for c in info.columns]
    assert table.row_count == info.row_count


def test_inspection_applies_column_types(tmp_path):
    path = messy_workbook(tmp_path / "budget.xlsx")
    insp = inspect_file(
        path,
        options=IngestOptions(
            sheet="Serials", header_row=1, column_types={"code": "string", "amount": "string"}
        ),
    )
    types = {c.name: c.inferred_type for c in insp.tables[0].columns}
    assert types["amount"] == "string" and types["code"] == "string"


def test_zip_bomb_part_is_refused(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/worksheets/sheet1.xml", "<a>" + "0" * 30_000_000 + "</a>")
    with pytest.raises(IngestError, match="zip bomb|expands"):
        inspect_file(path)


def test_large_sheet_is_streamed_and_preview_capped(tmp_path, monkeypatch):
    import analystos_engine.ingest as ingest_mod

    monkeypatch.setattr(ingest_mod, "INSPECT_MAX_ROWS", 100)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["id", "value"])
    for i in range(1000):
        ws.append([i, i * 2])
    path = tmp_path / "big.xlsx"
    wb.save(path)
    insp = inspect_file(path)
    assert insp.tables[0].row_count == 99
    assert any("larger than 100 rows" in w for w in insp.warnings)
    df, table, _ = ingest_mod.load_table_frame(path)
    assert len(df) == 1000
