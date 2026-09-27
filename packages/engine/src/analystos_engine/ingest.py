"""File inspection and ingestion into the workspace store.

Supported formats: CSV/TSV/TXT (encoding and dialect sniffed), Excel ``.xlsx``/``.xlsm``
(openpyxl, formulas read as cached values), JSON (array of records, ``{"data": [...]}``
wrappers, columnar objects), NDJSON / JSON Lines and Parquet.

Excel is treated as untrusted and messy:

* every sheet is scanned; hidden sheets are reported,
* blank rows and columns split a sheet into *table blocks* (several tables per sheet),
* title / note rows above a header (including merged title cells) are skipped,
* the header row is detected by scoring rows (text-heavy row followed by typed data),
* ``Total`` footer rows and trailing notes are excluded (and reported),
* formulas are read as their cached values (``data_only=True``); formula cells with no
  cached value are reported,
* Excel date serials in date-like columns are converted to dates,
* mixed-type columns are never silently coerced: a column is typed only when **every**
  non-null value conforms, otherwise it is kept as text and the inspection explains why.

Nothing here mutates the source file. The only write is into the workspace DuckDB
via :meth:`WorkspaceStore.write_table`.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import math
import re
import zipfile
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any, Literal

import numpy as np
import pandas as pd
from pydantic import BaseModel, ConfigDict, Field

from .store import StoreError, WorkspaceStore, validate_table_name
from .types import TableInfo, normalize_value

__all__ = [
    "IngestError",
    "CsvDialect",
    "DetectedColumn",
    "DetectedTable",
    "SkippedRow",
    "SheetInfo",
    "FileInspection",
    "IngestOptions",
    "inspect_file",
    "ingest_file",
    "load_table_frame",
    "normalize_column_name",
    "suggest_table_name",
    "detect_format",
    "DEFAULT_MAX_BYTES",
]

DEFAULT_MAX_BYTES = 500 * 1024 * 1024
PREVIEW_ROWS = 50
_HEADER_SCAN_ROWS = 25
_NA_TOKENS = frozenset(
    {
        "",
        "na",
        "n/a",
        "#n/a",
        "null",
        "none",
        "nan",
        "-",
        "--",
        "#null!",
        "#div/0!",
        "#value!",
        "#ref!",
        "#name?",
        "#num!",
        "nil",
    }
)
_TRUE = frozenset({"true", "yes", "y", "t"})
_FALSE = frozenset({"false", "no", "n", "f"})
_DATE_NAME_RE = re.compile(
    r"(date|day|month|period|_dt$|^dt_|time|created|updated|_at$|timestamp|week)", re.I
)
_TOTAL_RE = re.compile(r"^\s*(grand\s+)?(sub)?totals?\b", re.I)
_NUM_RE = re.compile(r"^\(?[+-]?\s*[$€£¥]?\s*(\d{1,3}(,\d{3})+|\d+)?(\.\d+)?([eE][+-]?\d+)?\s*%?\)?$")
_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y/%m/%d",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%d/%m/%Y",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y %H:%M:%S",
    "%d-%b-%Y",
    "%d-%b-%y",
    "%d %b %Y",
    "%b %d, %Y",
    "%B %d, %Y",
    "%d.%m.%Y",
    "%Y%m%d",
)
_EXCEL_EPOCH = dt.datetime(1899, 12, 30)

Format = Literal["csv", "excel", "json", "ndjson", "parquet"]


class IngestError(ValueError):
    """The file cannot be inspected or ingested."""


class _Model(BaseModel):
    model_config = ConfigDict(extra="forbid")


class CsvDialect(_Model):
    delimiter: str = ","
    quotechar: str = '"'
    encoding: str = "utf-8"
    has_bom: bool = False


class DetectedColumn(_Model):
    name: str
    original_name: str
    inferred_type: Literal["integer", "float", "boolean", "date", "timestamp", "string"]
    null_count: int = 0
    non_null_count: int = 0
    sample_values: list[Any] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)
    nonconforming_count: int = 0
    nonconforming_examples: list[str] = Field(default_factory=list)
    candidate_type: str | None = None


class SkippedRow(_Model):
    row: int
    reason: Literal["title", "blank", "total", "note", "malformed"]
    content: str = ""


class DetectedTable(_Model):
    key: str
    sheet: str | None = None
    name_suggestion: str
    header_row: int | None = None
    first_data_row: int
    last_data_row: int
    first_col: int = 1
    last_col: int = 1
    range: str | None = None
    columns: list[DetectedColumn] = Field(default_factory=list)
    row_count: int = 0
    preview_rows: list[list[Any]] = Field(default_factory=list)
    skipped_rows: list[SkippedRow] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)
    title: str | None = None


class SheetInfo(_Model):
    name: str
    visible: bool = True
    dimensions: str | None = None
    table_count: int = 0
    empty: bool = False
    max_row: int = 0
    max_col: int = 0


class FileInspection(_Model):
    path: str
    file_name: str
    format: Format
    size_bytes: int
    encoding: str | None = None
    dialect: CsvDialect | None = None
    sheets: list[SheetInfo] = Field(default_factory=list)
    tables: list[DetectedTable] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)

    def table(self, key: str | None = None) -> DetectedTable:
        if not self.tables:
            raise IngestError("no tables were detected in the file")
        if key is None:
            return self.tables[0]
        for t in self.tables:
            if t.key == key:
                return t
        for t in self.tables:
            if t.sheet == key:
                return t
        raise IngestError(f"no detected table {key!r}; available: {[t.key for t in self.tables]}")


class IngestOptions(_Model):
    table_name: str | None = None
    table_key: str | None = None
    sheet: str | None = None
    header_row: int | None = Field(
        default=None, ge=1, description="1-based row of the header (overrides detection)"
    )
    if_exists: Literal["fail", "replace", "append"] = "fail"
    column_types: dict[str, Literal["integer", "float", "boolean", "date", "timestamp", "string"]] = Field(
        default_factory=dict
    )
    column_names: dict[str, str] = Field(
        default_factory=dict, description="rename map: detected name -> new name"
    )
    max_bytes: int = Field(default=DEFAULT_MAX_BYTES, gt=0)
    delimiter: str | None = None
    encoding: str | None = None


# ------------------------------------------------------------------------ naming


def normalize_column_name(name: Any, index: int) -> str:
    text = "" if name is None else str(name)
    text = text.strip().lower()
    text = re.sub(r"\(.*?\)|\[.*?\]", "", text)
    text = text.replace("%", " pct ").replace("#", " num ").replace("&", " and ")
    text = re.sub(r"[^0-9a-z]+", "_", text).strip("_")
    text = re.sub(r"_+", "_", text)
    if not text:
        return f"column_{index + 1}"
    if text[0].isdigit():
        text = f"c_{text}"
    return text[:63]


def _dedupe(names: list[str]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for n in names:
        if n in seen:
            seen[n] += 1
            cand = f"{n}_{seen[n]}"
            while cand in seen:
                seen[n] += 1
                cand = f"{n}_{seen[n]}"
            seen[cand] = 1
            out.append(cand)
        else:
            seen[n] = 1
            out.append(n)
    return out


def suggest_table_name(*parts: str | None) -> str:
    raw = "_".join(p for p in parts if p)
    name = normalize_column_name(raw, 0)
    if name.startswith("column_") and not raw:
        name = "dataset"
    if name.startswith("_aos_"):
        name = "t" + name
    return name[:63]


# ------------------------------------------------------------------------ format


def detect_format(path: Path) -> Format:
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv", ".txt", ".tab"):
        return "csv"
    if ext in (".xlsx", ".xlsm"):
        return "excel"
    if ext == ".xls":
        raise IngestError("legacy .xls workbooks are not supported; save the file as .xlsx and upload again")
    if ext == ".json":
        with path.open("rb") as fh:
            head = fh.read(4096).lstrip()
        if head.startswith(b"{"):
            # NDJSON when the first line is a complete object and another follows
            first, _, rest = head.partition(b"\n")
            if rest.strip().startswith(b"{"):
                try:
                    json.loads(first)
                    return "ndjson"
                except ValueError:
                    pass
        return "json"
    if ext in (".jsonl", ".ndjson"):
        return "ndjson"
    if ext in (".parquet", ".pq"):
        return "parquet"
    raise IngestError(
        f"unsupported file type {ext or '(none)'}; supported: csv, tsv, txt, xlsx, xlsm, json, ndjson, parquet"
    )


def _check_file(path: str | Path, max_bytes: int) -> Path:
    p = Path(path)
    if not p.exists():
        raise IngestError(f"file not found: {p.name}")
    if not p.is_file():
        raise IngestError(f"not a regular file: {p.name}")
    size = p.stat().st_size
    if size == 0:
        raise IngestError(f"file {p.name} is empty")
    if size > max_bytes:
        raise IngestError(f"file {p.name} is {size:,} bytes; the limit is {max_bytes:,} bytes")
    return p


# ------------------------------------------------------------------- type inference


def _is_na(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str) and v.strip().lower() in _NA_TOKENS:
        return True
    return v is pd.NaT


def _num_to_text(v: Any) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _to_text(v: Any) -> str:
    if isinstance(v, dt.datetime):
        return (
            v.isoformat(sep=" ")
            if (v.hour, v.minute, v.second, v.microsecond) != (0, 0, 0, 0)
            else v.date().isoformat()
        )
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, int | float | bool | np.integer | np.floating):
        return _num_to_text(v.item() if hasattr(v, "item") else v)
    return str(v).strip()


def _parse_numbers(s: pd.Series) -> tuple[pd.Series, list[str]]:
    """Vectorised numeric parse of a string series (NA already removed)."""
    notes: list[str] = []
    text = s.str.strip()
    has_pct = text.str.endswith("%")
    has_cur = text.str.contains(r"[$€£¥]", regex=True)
    has_thousands = text.str.contains(r"\d,\d{3}", regex=True)
    paren = text.str.startswith("(") & text.str.endswith(")")
    cleaned = text.str.replace(r"[$€£¥,%\s()]", "", regex=True)
    nums = pd.to_numeric(cleaned, errors="coerce")
    valid_shape = text.str.match(_NUM_RE)
    nums = nums.where(valid_shape)
    nums = nums.where(~paren, -nums.abs())
    if bool(has_pct.any()):
        if bool(has_pct.all()):
            nums = nums / 100.0
            notes.append("percent strings converted to fractions (12% -> 0.12)")
        else:
            nums = nums.where(~has_pct)  # mixed percent/non-percent is ambiguous: do not convert
    if bool(has_cur.any()):
        notes.append("currency symbols removed")
    if bool(has_thousands.any()):
        notes.append("thousands separators removed")
    if bool(paren.any()):
        notes.append("accounting negatives in parentheses converted")
    return nums, notes


def _best_date_format(sample: list[str]) -> str | None:
    best, best_hits = None, 0
    for fmt in _DATE_FORMATS:
        hits = 0
        for v in sample:
            try:
                dt.datetime.strptime(v, fmt)
                hits += 1
            except ValueError:
                continue
        if hits > best_hits:
            best, best_hits = fmt, hits
            if hits == len(sample):
                break
    if best is None or best_hits < max(1, int(0.5 * len(sample))):
        return None
    return best


def _parse_dates(s: pd.Series) -> tuple[pd.Series, str | None]:
    sample = s.dropna().astype(str).head(200).tolist()
    sample = [v.strip() for v in sample]
    if not sample or not any(ch.isdigit() for ch in "".join(sample[:20])):
        return pd.Series(pd.NaT, index=s.index), None
    if all(re.fullmatch(r"\d+", v) for v in sample) and not all(
        re.fullmatch(r"(19|20)\d{6}", v) for v in sample
    ):
        return pd.Series(pd.NaT, index=s.index), None
    fmt = _best_date_format(sample)
    if fmt is None:
        return pd.Series(pd.NaT, index=s.index), None
    parsed = pd.to_datetime(s.str.strip(), format=fmt, errors="coerce")
    return parsed, fmt


def _dates_are_midnight(parsed: pd.Series) -> bool:
    valid = parsed.dropna()
    if valid.empty:
        return True
    return bool(
        (
            (valid.dt.hour == 0)
            & (valid.dt.minute == 0)
            & (valid.dt.second == 0)
            & (valid.dt.microsecond == 0)
        ).all()
    )


def _col_result(
    name: str,
    original: str,
    typ: str,
    series: pd.Series,
    null_count: int,
    notes: list[str],
    nonconf: list[str] | None = None,
    nonconf_count: int = 0,
    candidate: str | None = None,
) -> tuple[pd.Series, DetectedColumn]:
    samples = [normalize_value(v) for v in series.dropna().head(5).tolist()]
    return (
        series,
        DetectedColumn(
            name=name,
            original_name=original,
            inferred_type=typ,  # type: ignore[arg-type]
            null_count=null_count,
            non_null_count=int(series.notna().sum()),
            sample_values=samples,
            notes=notes,
            nonconforming_count=nonconf_count,
            nonconforming_examples=(nonconf or [])[:5],
            candidate_type=candidate,
        ),
    )


def infer_column(
    values: Sequence[Any] | pd.Series, name: str, original: str | None = None, *, force: str | None = None
) -> tuple[pd.Series, DetectedColumn]:
    """Infer the type of a column of raw values and return the typed series.

    A column is typed (number/date/boolean) only when every non-null value conforms.
    Otherwise it stays text and ``candidate_type``/``nonconforming_*`` explain what
    blocked the conversion. ``force`` converts anyway (non-conforming -> NULL, counted).
    """
    original = original if original is not None else name
    raw = (
        pd.Series(list(values), dtype=object) if not isinstance(values, pd.Series) else values.astype(object)
    )
    raw = raw.reset_index(drop=True)
    na_mask = raw.map(_is_na).astype(bool)
    null_count = int(na_mask.sum())
    present = raw[~na_mask]
    notes: list[str] = []
    if present.empty:
        typ = force or "string"
        return _col_result(name, original, typ, _empty_series(len(raw), typ), null_count, ["column is empty"])

    kinds = present.map(_kind)
    kind_set = set(kinds.unique())

    # ---- all native numbers (Excel numeric cells, JSON numbers)
    if kind_set <= {"int", "float"} and force in (None, "integer", "float", "date", "timestamp"):
        nums = pd.to_numeric(present, errors="coerce").astype(float)
        if force in ("date", "timestamp") or (force is None and _looks_like_serial(name, nums)):
            conv = _serial_to_datetime(nums)
            full = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
            full[~na_mask] = conv.values
            typ = "date" if _dates_are_midnight(full) else "timestamp"
            notes.append(f"converted {len(nums)} Excel date serial numbers to dates")
            out = full.dt.date if typ == "date" else full
            return _col_result(name, original, typ, _as_obj_dates(out, typ), null_count, notes)
        return _numeric_result(name, original, raw, na_mask, nums, null_count, notes, force)

    # ---- all native datetimes
    if kind_set <= {"datetime", "date"} and force in (None, "date", "timestamp"):
        full = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
        full[~na_mask] = pd.to_datetime(
            present.map(lambda v: v if isinstance(v, dt.datetime) else dt.datetime.combine(v, dt.time()))
        )
        typ = force or ("date" if _dates_are_midnight(full) else "timestamp")
        return _col_result(
            name,
            original,
            typ,
            _as_obj_dates(full.dt.date if typ == "date" else full, typ),
            null_count,
            notes,
        )

    if kind_set == {"bool"} and force in (None, "boolean"):
        out = pd.Series([None] * len(raw), dtype=object)
        out[~na_mask] = present.astype(bool).values
        return _col_result(name, original, "boolean", out.astype("boolean"), null_count, notes)

    # ---- text path (also for mixed native types): stringify then parse
    text = present.map(_to_text)
    mixed_native = len(kind_set - {"str"}) > 0 and "str" in kind_set
    if mixed_native:
        notes.append("column mixes cell types (" + ", ".join(sorted(kind_set)) + ")")

    if force == "string":
        out = pd.Series([None] * len(raw), dtype=object)
        out[~na_mask] = text.values
        return _col_result(name, original, "string", out, null_count, notes)

    lowered = text.str.lower()
    if (
        force in (None, "boolean")
        and lowered.isin(_TRUE | _FALSE).all()
        and (force == "boolean" or not lowered.isin({"y", "n", "t", "f"}).all())
    ):
        out = pd.Series([None] * len(raw), dtype=object)
        out[~na_mask] = lowered.isin(_TRUE).values
        return _col_result(name, original, "boolean", out.astype("boolean"), null_count, notes)

    # leading-zero identifiers (zip codes, SKUs) stay text
    leading_zero = text.str.match(r"^0\d+$")
    if force is None and bool(leading_zero.any()):
        out = pd.Series([None] * len(raw), dtype=object)
        out[~na_mask] = text.values
        notes.append("values with leading zeros kept as text (identifiers)")
        return _col_result(name, original, "string", out, null_count, notes)

    if (
        force in (None, "date")
        and _DATE_NAME_RE.search(name or "")
        and bool(text.str.fullmatch(r"(19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])").all())
    ):
        parsed = pd.to_datetime(text, format="%Y%m%d", errors="coerce")
        if bool(parsed.notna().all()):
            full = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
            full[~na_mask] = parsed.values
            notes.append("yyyymmdd numbers read as dates")
            return _col_result(name, original, "date", _as_obj_dates(full.dt.date, "date"), null_count, notes)

    nums, num_notes = _parse_numbers(text)
    num_ok = nums.notna()
    if force in ("integer", "float") or (force is None and bool(num_ok.all())):
        if (
            force is None
            and _looks_like_serial(name, nums)
            and bool(num_ok.all())
            and not text.str.contains(r"[.,$%]").any()
        ):
            conv = _serial_to_datetime(nums)
            full = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
            full[~na_mask] = conv.values
            typ = "date" if _dates_are_midnight(full) else "timestamp"
            notes.append(f"converted {len(nums)} Excel date serial numbers to dates")
            return _col_result(
                name,
                original,
                typ,
                _as_obj_dates(full.dt.date if typ == "date" else full, typ),
                null_count,
                notes,
            )
        bad = text[~num_ok]
        extra = []
        if len(bad):
            extra.append(f"{len(bad)} values could not be converted to numbers and were set to NULL")
        return _numeric_result(
            name,
            original,
            raw,
            na_mask,
            nums.astype(float),
            null_count,
            notes + num_notes + extra,
            force,
            nonconf=bad.tolist(),
        )

    parsed, fmt = _parse_dates(text)
    date_ok = parsed.notna()
    if force in ("date", "timestamp") or (force is None and fmt is not None and bool(date_ok.all())):
        if fmt is None and force:
            parsed = pd.to_datetime(text, errors="coerce", format="mixed")
            date_ok = parsed.notna()
        full = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
        full[~na_mask] = parsed.values
        typ = force or ("date" if _dates_are_midnight(full) else "timestamp")
        if fmt:
            notes.append(f"dates parsed with format {fmt}")
        bad = text[~date_ok]
        if len(bad):
            notes.append(f"{len(bad)} values could not be parsed as dates and were set to NULL")
        return _col_result(
            name,
            original,
            typ,
            _as_obj_dates(full.dt.date if typ == "date" else full, typ),
            null_count,
            notes,
            bad.tolist(),
            len(bad),
        )

    # ---- stays text; explain near-misses so the analyst can decide
    out = pd.Series([None] * len(raw), dtype=object)
    out[~na_mask] = text.values
    candidate = None
    nonconf: list[str] = []
    ratio_num = float(num_ok.mean())
    ratio_date = float(date_ok.mean()) if fmt else 0.0
    if ratio_num >= 0.5 and ratio_num >= ratio_date:
        candidate = "float"
        nonconf = text[~num_ok].tolist()
        notes.append(
            f"mixed types: {int(num_ok.sum())} numeric and {len(nonconf)} non-numeric values; kept as text "
            "(override the type to convert, non-numeric values become NULL)"
        )
    elif ratio_date >= 0.5:
        candidate = "date"
        nonconf = text[~date_ok].tolist()
        notes.append(
            f"{len(nonconf)} values do not parse as dates ({fmt}); kept as text "
            "(override the type to convert, malformed dates become NULL)"
        )
    return _col_result(name, original, "string", out, null_count, notes, nonconf, len(nonconf), candidate)


def _numeric_result(
    name: str,
    original: str,
    raw: pd.Series,
    na_mask: pd.Series,
    nums: pd.Series,
    null_count: int,
    notes: list[str],
    force: str | None,
    nonconf: list[str] | None = None,
) -> tuple[pd.Series, DetectedColumn]:
    valid = nums.dropna()
    integral = bool((valid == np.floor(valid)).all()) and bool((valid.abs() < 2**53).all())
    typ = force if force in ("integer", "float") else ("integer" if integral else "float")
    if typ == "integer" and not integral:
        notes.append("fractional values rounded to integers")
    full = pd.Series([None] * len(raw), dtype=object)
    full[~na_mask] = nums.values
    if typ == "integer":
        out = pd.array(
            [None if v is None or (isinstance(v, float) and math.isnan(v)) else int(round(v)) for v in full],
            dtype="Int64",
        )
        series = pd.Series(out)
    else:
        series = pd.to_numeric(full, errors="coerce").astype("float64")
    return _col_result(name, original, typ, series, null_count, notes, nonconf, len(nonconf or []))


def _kind(v: Any) -> str:
    if isinstance(v, bool | np.bool_):
        return "bool"
    if isinstance(v, int | np.integer):
        return "int"
    if isinstance(v, float | np.floating):
        return "float"
    if isinstance(v, dt.datetime | pd.Timestamp):
        return "datetime"
    if isinstance(v, dt.date):
        return "date"
    if isinstance(v, dt.time):
        return "time"
    return "str"


def _looks_like_serial(name: str, nums: pd.Series) -> bool:
    valid = nums.dropna()
    if valid.empty or not _DATE_NAME_RE.search(name or ""):
        return False
    return bool(((valid >= 20000) & (valid <= 80000)).all())


def _serial_to_datetime(nums: pd.Series) -> pd.Series:
    return pd.to_datetime(_EXCEL_EPOCH) + pd.to_timedelta(nums.astype(float), unit="D").dt.round("s")


def _as_obj_dates(s: pd.Series, typ: str) -> pd.Series:
    if typ == "date":
        return pd.Series(
            [
                None if (v is None or v is pd.NaT or (isinstance(v, float) and math.isnan(v))) else v
                for v in s
            ],
            dtype=object,
        )
    return pd.to_datetime(s)


def _empty_series(n: int, typ: str) -> pd.Series:
    if typ == "integer":
        return pd.Series(pd.array([None] * n, dtype="Int64"))
    if typ == "float":
        return pd.Series([np.nan] * n, dtype="float64")
    if typ == "boolean":
        return pd.Series(pd.array([None] * n, dtype="boolean"))
    if typ == "timestamp":
        return pd.Series(pd.NaT, index=range(n), dtype="datetime64[ns]")
    return pd.Series([None] * n, dtype=object)


# ------------------------------------------------------------------- grid analysis


def _clean_cell(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def _col_letter(idx: int) -> str:
    s = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def _is_texty(v: Any) -> bool:
    return isinstance(v, str) and not _NUM_RE.match(v) and v.strip().lower() not in _NA_TOKENS


def _typed(v: Any) -> bool:
    if v is None:
        return False
    if isinstance(v, int | float | dt.date | dt.datetime | bool):
        return True
    return isinstance(v, str) and bool(_NUM_RE.match(v)) and any(ch.isdigit() for ch in v)


class _Block:
    def __init__(self, r0: int, r1: int, c0: int, c1: int) -> None:
        self.r0, self.r1, self.c0, self.c1 = r0, r1, c0, c1  # inclusive, 0-based


def _components(occupied: list[list[bool]]) -> list[_Block]:
    """Bounding boxes of 8-connected groups of occupied cells."""
    h = len(occupied)
    w = max((len(r) for r in occupied), default=0)
    seen = [[False] * w for _ in range(h)]
    blocks: list[_Block] = []
    for r in range(h):
        row = occupied[r]
        for c in range(len(row)):
            if not row[c] or seen[r][c]:
                continue
            stack = [(r, c)]
            seen[r][c] = True
            r0 = r1 = r
            c0 = c1 = c
            while stack:
                y, x = stack.pop()
                r0, r1, c0, c1 = min(r0, y), max(r1, y), min(c0, x), max(c1, x)
                for dy in (-1, 0, 1):
                    ny = y + dy
                    if ny < 0 or ny >= h:
                        continue
                    orow = occupied[ny]
                    for dx in (-1, 0, 1):
                        nx = x + dx
                        if 0 <= nx < len(orow) and orow[nx] and not seen[ny][nx]:
                            seen[ny][nx] = True
                            stack.append((ny, nx))
            blocks.append(_Block(r0, r1, c0, c1))
    blocks.sort(key=lambda b: (b.r0, b.c0))
    return blocks


def _row_signature(grid: list[list[Any]], r: int, c0: int, c1: int) -> list[str]:
    out = []
    for c in range(c0, c1 + 1):
        v = grid[r][c] if c < len(grid[r]) else None
        out.append("empty" if v is None else ("text" if _is_texty(v) else "typed"))
    return out


def _is_note_block(grid: list[list[Any]], b: _Block) -> bool:
    """Every row holds at most one value: a title / notes block, not a table."""
    for r in range(b.r0, b.r1 + 1):
        if sum(1 for c in range(b.c0, b.c1 + 1) if c < len(grid[r]) and grid[r][c] is not None) > 1:
            return False
    return (b.c1 - b.c0) >= 1 or (b.r1 - b.r0) <= 2


def _continues(grid: list[list[Any]], a: _Block, b: _Block) -> bool:
    """Does block ``b`` (below ``a``) continue ``a``'s data after blank rows?"""
    gap = b.r0 - a.r1 - 1
    if gap < 1 or gap > 3:
        return False
    lo, hi = max(a.c0, b.c0), min(a.c1, b.c1)
    if hi < lo:
        return False
    narrow = min(a.c1 - a.c0, b.c1 - b.c0) + 1
    if (hi - lo + 1) < 0.5 * narrow or b.c0 < a.c0 - 1 or b.c1 > a.c1 + 1:
        return False
    first = [grid[b.r0][c] if c < len(grid[b.r0]) else None for c in range(b.c0, b.c1 + 1)]
    filled = [v for v in first if v is not None]
    if filled and isinstance(filled[0], str) and _TOTAL_RE.match(filled[0]):
        return True
    if a.r1 - a.r0 < 1:
        return False  # a single-row block has no data rows to compare with
    sig_b = _row_signature(grid, b.r0, a.c0, a.c1)
    sig_a = _row_signature(grid, a.r1, a.c0, a.c1)
    typed_b = [i for i, s in enumerate(sig_b) if s == "typed"]
    if not typed_b:
        return False
    return all(sig_a[i] in ("typed", "empty") for i in typed_b)


def _find_blocks(
    grid: list[list[Any]], merged_ranges: list[tuple[int, int, int, int]] | None = None
) -> list[_Block]:
    occupied = [[v is not None for v in row] for row in grid]
    for r0, r1, c0, c1 in merged_ranges or []:
        for r in range(r0, min(r1 + 1, len(occupied))):
            for c in range(c0, min(c1 + 1, len(occupied[r]))):
                occupied[r][c] = True
    blocks = _components(occupied)
    merged = True
    while merged:
        merged = False
        for i, a in enumerate(blocks):
            for j, b in enumerate(blocks):
                if i == j or b.r0 <= a.r1:
                    continue
                if _continues(grid, a, b):
                    a.r1 = max(a.r1, b.r1)
                    a.c0, a.c1 = min(a.c0, b.c0), max(a.c1, b.c1)
                    del blocks[j]
                    merged = True
                    break
            if merged:
                break
    return blocks


def _is_year(v: Any) -> bool:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return False
    return f.is_integer() and 1900 <= f <= 2100


def _header_score(row: list[Any], nxt: list[list[Any]], width: int) -> float:
    filled = [v for v in row if v is not None]
    if not filled:
        return -1.0
    fill = len(filled) / width
    texty = sum(1 for v in filled if _is_texty(v)) / len(filled)
    unique = len({str(v).lower() for v in filled}) / len(filled)
    typed_below = 0.0
    if nxt:
        typed_cells = sum(1 for r in nxt for v in r if _typed(v))
        all_cells = sum(1 for r in nxt for v in r if v is not None)
        typed_below = typed_cells / all_cells if all_cells else 0.0
    return fill * 2 + texty * 2 + unique + typed_below


def _analyse_block(
    grid: list[list[Any]],
    b: _Block,
    *,
    header_override: int | None = None,
) -> tuple[int | None, int, int, list[SkippedRow], str | None]:
    """Return (header_row_idx, first_data_idx, last_data_idx, skipped, title) for a block (0-based rows)."""
    width = b.c1 - b.c0 + 1
    rows = {
        r: [grid[r][c] if c < len(grid[r]) else None for c in range(b.c0, b.c1 + 1)]
        for r in range(b.r0, b.r1 + 1)
    }
    skipped: list[SkippedRow] = []
    title: str | None = None
    h: int | None
    if header_override is not None:
        h = header_override
        for r in range(b.r0, h):
            vals = [str(v) for v in rows.get(r, []) if v is not None]
            if vals:
                skipped.append(SkippedRow(row=r + 1, reason="title", content=" ".join(vals)[:200]))
        first = h + 1
    else:
        candidates = list(range(b.r0, min(b.r1 + 1, b.r0 + _HEADER_SCAN_ROWS)))
        best, best_score = None, -1.0
        for r in candidates:
            row = rows[r]
            filled = [v for v in row if v is not None]
            if width >= 2 and len(filled) == 1 and isinstance(filled[0], str):
                continue  # title or note row
            n_text = sum(1 for v in filled if _is_texty(v))
            non_text_ok = all(_is_texty(v) or _is_year(v) for v in filled)
            if n_text == 0 or not (non_text_ok or n_text / len(filled) >= 0.75):
                continue
            below = [rows[x] for x in range(r + 1, min(b.r1 + 1, r + 6))]
            score = _header_score(row, below, width) - 0.1 * (r - b.r0)
            if score > best_score:
                best, best_score = r, score
        h = best if best is not None and (best < b.r1 or best == b.r0 == b.r1) else None
        start = h if h is not None else b.r0
        for r in range(b.r0, start):
            vals = [str(v) for v in rows[r] if v is not None]
            skipped.append(SkippedRow(row=r + 1, reason="title", content=" ".join(vals)[:200]))
            if title is None and vals:
                title = " ".join(vals)[:200]
        first = (h + 1) if h is not None else b.r0
    last = b.r1
    # trailing totals / notes
    while last >= first:
        row = rows.get(last, [])
        filled = [v for v in row if v is not None]
        if filled and isinstance(filled[0], str) and _TOTAL_RE.match(filled[0]) and row.index(filled[0]) <= 1:
            skipped.append(SkippedRow(row=last + 1, reason="total", content=" ".join(map(str, filled))[:200]))
            last -= 1
            continue
        if (
            width >= 3
            and len(filled) == 1
            and isinstance(filled[0], str)
            and len(filled[0]) > 12
            and last > first
        ):
            skipped.append(SkippedRow(row=last + 1, reason="note", content=filled[0][:200]))
            last -= 1
            continue
        break
    return (h if header_override is None else header_override), first, last, skipped, title


def _table_from_grid(
    grid: list[list[Any]],
    b: _Block,
    header_idx: int | None,
    first: int,
    last: int,
) -> tuple[list[str], list[list[Any]]]:
    width = b.c1 - b.c0 + 1
    header = [
        grid[header_idx][c] if header_idx is not None and c < len(grid[header_idx]) else None
        for c in range(b.c0, b.c1 + 1)
    ]
    data = []
    for r in range(first, last + 1):
        row = [grid[r][c] if c < len(grid[r]) else None for c in range(b.c0, b.c1 + 1)]
        data.append(row)
    names = [str(h) if h is not None else "" for h in header] if header_idx is not None else [""] * width
    return names, data


def _frame_from_rows(
    original_names: list[str],
    rows: list[list[Any]],
    *,
    column_types: Mapping[str, str] | None = None,
) -> tuple[pd.DataFrame, list[DetectedColumn], list[str]]:
    """Build a typed DataFrame; drops fully empty columns (reported)."""
    notes: list[str] = []
    width = len(original_names)
    names = _dedupe([normalize_column_name(n, i) for i, n in enumerate(original_names)])
    blank_rows = [i for i, r in enumerate(rows) if all(_is_na(v) for v in r)]
    if blank_rows:
        rows = [r for i, r in enumerate(rows) if i not in set(blank_rows)]
        notes.append(f"{len(blank_rows)} blank rows removed")
    cols: dict[str, pd.Series] = {}
    detected: list[DetectedColumn] = []
    dropped: list[str] = []
    for i in range(width):
        values = [r[i] if i < len(r) else None for r in rows]
        if all(_is_na(v) for v in values) and not original_names[i]:
            dropped.append(names[i])
            continue
        force = (column_types or {}).get(names[i])
        series, info = infer_column(values, names[i], original_names[i] or names[i], force=force)
        cols[names[i]] = series.reset_index(drop=True)
        detected.append(info)
    if dropped:
        notes.append(f"{len(dropped)} empty unnamed columns removed")
    df = pd.DataFrame(cols) if cols else pd.DataFrame()
    return df, detected, notes


# ------------------------------------------------------------------------ Excel


XLSX_MAX_EXPANDED = 500 * 1024 * 1024
"""Absolute cap on the uncompressed size of all parts of a workbook."""
XLSX_MAX_RATIO = 100
"""Compression ratio above which a large part is treated as a zip bomb."""
INSPECT_MAX_ROWS = 20_000
"""Rows read per sheet when inspecting (previewing) a workbook; ingest reads everything."""


def _check_zip(path: Path, max_bytes: int) -> None:
    """Refuse workbooks whose parts expand far beyond the upload size before parsing them.

    Every part must stay under ``XLSX_MAX_RATIO``:1 once it is over 1 MB uncompressed, and
    the whole workbook may expand to at most the smaller of ``XLSX_MAX_EXPANDED`` and
    ``max(20 x max_bytes, 50 MB)``.
    """
    try:
        with zipfile.ZipFile(path) as z:
            total = 0
            for info in z.infolist():
                total += info.file_size
                if (
                    info.file_size > 1_000_000
                    and info.file_size / max(info.compress_size, 1) > XLSX_MAX_RATIO
                ):
                    raise IngestError(
                        f"workbook part {info.filename} expands {info.file_size / max(info.compress_size, 1):,.0f}x; "
                        "refusing a likely zip bomb"
                    )
            limit = min(XLSX_MAX_EXPANDED, max(max_bytes * 20, 50 * 1024 * 1024))
            if total > limit:
                raise IngestError(f"workbook expands to {total:,} bytes, above the {limit:,} byte limit")
    except zipfile.BadZipFile as exc:
        raise IngestError("file is not a valid .xlsx workbook") from exc


def _load_workbook(path: Path) -> Any:
    """Open a workbook in streaming (read-only) mode: rows are parsed as they are read, so
    memory stays proportional to what is kept, not to the whole sheet XML tree."""
    import openpyxl

    try:
        return openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
    except Exception as exc:  # openpyxl raises many exception types for corrupt files
        raise IngestError(f"cannot read workbook: {type(exc).__name__}: {exc}") from exc


_MERGE_RE = re.compile(rb'<mergeCell\s+ref="([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?"')


def _col_index(letters: bytes) -> int:
    n = 0
    for ch in letters.decode():
        n = n * 26 + (ord(ch) - 64)
    return n


def _merged_ranges(path: Path, part: str | None) -> list[tuple[int, int, int, int]]:
    """Merged cell ranges of one sheet, read from its XML (read-only sheets do not expose them)."""
    if not part:
        return []
    out: list[tuple[int, int, int, int]] = []
    try:
        with zipfile.ZipFile(path) as z, z.open(part.lstrip("/")) as fh:
            tail = b""
            while chunk := fh.read(1 << 20):
                buf = tail + chunk
                cut = buf.rfind(b"<")
                body, tail = (buf[:cut], buf[cut:]) if cut > 0 else (buf, b"")
                for m in _MERGE_RE.finditer(body):
                    c0, r0 = _col_index(m.group(1)), int(m.group(2))
                    c1 = _col_index(m.group(3)) if m.group(3) else c0
                    r1 = int(m.group(4)) if m.group(4) else r0
                    out.append((r0 - 1, r1 - 1, c0 - 1, c1 - 1))
    except (KeyError, zipfile.BadZipFile):
        return out
    return out


def _formula_cells_without_values(path: Path) -> dict[str, int]:
    """Count formula cells per sheet whose cached value is missing (never calculated)."""
    out: dict[str, int] = {}
    try:
        with zipfile.ZipFile(path) as z:
            import openpyxl

            wb = openpyxl.load_workbook(path, data_only=False, read_only=True, keep_links=False)
            names = wb.sheetnames
            wb.close()
            sheet_files = sorted(n for n in z.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n))
            for idx, fname in enumerate(sorted(sheet_files, key=lambda n: int(re.findall(r"\d+", n)[-1]))):
                xml = z.read(fname).decode("utf-8", errors="replace")
                missing = len(re.findall(r"<f>[^<]*</f>\s*(?:<v\s*/>|<v></v>|(?=</c>))", xml))
                missing += len(re.findall(r"<f\s[^>]*/>\s*(?:<v\s*/>|<v></v>|(?=</c>))", xml))
                if missing and idx < len(names):
                    out[names[idx]] = missing
    except Exception:  # best effort diagnostics only
        return out
    return out


def _sheet_grid(
    ws: Any, path: Path, max_rows: int | None = None
) -> tuple[list[list[Any]], list[tuple[int, int, int, int]], bool]:
    """Cell values of a sheet (at most ``max_rows`` rows), its merged ranges and whether the
    row cap cut the sheet short."""
    grid: list[list[Any]] = []
    truncated = False
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()  # read to the real end even when the stored dimension is wrong
    for row in ws.iter_rows(values_only=True):
        if max_rows is not None and len(grid) >= max_rows:
            truncated = True
            break
        grid.append([_clean_cell(v) for v in row])
    merged = _merged_ranges(path, getattr(ws, "_worksheet_path", None))
    merged_cells = getattr(ws, "merged_cells", None)
    for rng in merged_cells.ranges if merged_cells is not None else []:
        merged.append((rng.min_row - 1, rng.max_row - 1, rng.min_col - 1, rng.max_col - 1))
    width = 0
    for row in grid:
        for i in range(len(row) - 1, -1, -1):
            if row[i] is not None:
                width = max(width, i + 1)
                break
    grid = [row[:width] + [None] * (width - len(row[:width])) for row in grid]
    while grid and all(v is None for v in grid[-1]):
        grid.pop()
    return grid, merged, truncated


def _block_text(grid: list[list[Any]], b: _Block) -> str:
    parts = []
    for r in range(b.r0, b.r1 + 1):
        vals = [str(grid[r][c]) for c in range(b.c0, b.c1 + 1) if c < len(grid[r]) and grid[r][c] is not None]
        if vals:
            parts.append(" ".join(vals))
    return " / ".join(parts)[:300]


def _excel_tables(
    path: Path,
    *,
    sheet_filter: str | None = None,
    header_override: int | None = None,
    max_rows: int | None = None,
) -> tuple[list[SheetInfo], list[tuple[DetectedTable, list[str], list[list[Any]]]], list[str]]:
    wb = _load_workbook(path)
    warnings: list[str] = []
    missing = _formula_cells_without_values(path)
    sheets: list[SheetInfo] = []
    tables: list[tuple[DetectedTable, list[str], list[list[Any]]]] = []
    stem = path.stem
    try:
        for ws in wb.worksheets:
            if sheet_filter is not None and ws.title != sheet_filter:
                continue
            grid, merged, truncated = _sheet_grid(ws, path, max_rows)
            visible = getattr(ws, "sheet_state", "visible") == "visible"
            if truncated:
                warnings.append(
                    f"sheet {ws.title!r} is larger than {max_rows:,} rows; the preview reads the first {max_rows:,} "
                    "rows and ingest reads all of them"
                )
            info = SheetInfo(
                name=ws.title,
                visible=visible,
                dimensions=f"A1:{_col_letter(len(grid[0]))}{len(grid)}" if grid and grid[0] else None,
                max_row=len(grid),
                max_col=len(grid[0]) if grid else 0,
                empty=not grid,
            )
            if not visible:
                warnings.append(f"sheet {ws.title!r} is hidden")
            if ws.title in missing:
                warnings.append(
                    f"sheet {ws.title!r} has {missing[ws.title]} formula cells without cached values "
                    "(the workbook was never recalculated in Excel); those cells are read as empty"
                )
            if header_override is not None and grid:
                h = header_override - 1
                if h >= len(grid):
                    raise IngestError(
                        f"header row {header_override} is beyond the last row of sheet {ws.title!r}"
                    )
                below = [r for r in range(h + 1, len(grid)) if any(v is not None for v in grid[r])]
                filled_cols = [c for c in range(len(grid[h])) if grid[h][c] is not None] or [0]
                c0 = c1 = filled_cols[0]
                while c1 + 1 < len(grid[h]) and grid[h][c1 + 1] is not None:
                    c1 += 1  # the header's contiguous run of cells defines the table width
                blocks = [_Block(0, below[-1] if below else h, c0, c1)]
            else:
                blocks = _find_blocks(grid, merged) if grid else []
            generic = re.fullmatch(r"sheet\s*\d*", ws.title.strip(), re.I) is not None
            pending_title: dict[int, str] = {}
            n_tables = 0
            for b in blocks:
                if header_override is None and _is_note_block(grid, b):
                    for other in blocks:
                        if other is not b and other.r0 > b.r1 and other.c0 <= b.c1 and other.c1 >= b.c0:
                            pending_title.setdefault(id(other), _block_text(grid, b))
                            break
                    continue
                h_override = None if header_override is None else header_override - 1
                header_idx, first, last, skipped, title = _analyse_block(grid, b, header_override=h_override)
                if last < first:
                    continue
                n_tables += 1
                names, data = _table_from_grid(grid, b, header_idx, first, last)
                top = (header_idx if header_idx is not None else first) + 1
                key = f"{ws.title}!{_col_letter(b.c0 + 1)}{top}:{_col_letter(b.c1 + 1)}{last + 1}"
                base = stem if generic else ws.title
                name = suggest_table_name(base) if n_tables == 1 else suggest_table_name(base, str(n_tables))
                title = pending_title.get(id(b)) or title
                table = DetectedTable(
                    key=key,
                    sheet=ws.title,
                    name_suggestion=name,
                    header_row=header_idx + 1 if header_idx is not None else None,
                    first_data_row=first + 1,
                    last_data_row=last + 1,
                    first_col=b.c0 + 1,
                    last_col=b.c1 + 1,
                    range=key.split("!", 1)[1],
                    skipped_rows=skipped,
                    title=title,
                )
                if header_idx is None:
                    table.notes.append("no header row detected; columns named column_1, column_2, ...")
                tables.append((table, names, data))
            info.table_count = n_tables
            sheets.append(info)
    finally:
        wb.close()
    if sheet_filter is not None and not sheets:
        raise IngestError(f"sheet {sheet_filter!r} not found")
    return sheets, tables, warnings


# -------------------------------------------------------------------------- CSV


def _sniff_encoding(sample: bytes) -> tuple[str, bool]:
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", True
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16", True
    try:
        sample.decode("utf-8")
        return "utf-8", False
    except UnicodeDecodeError as exc:
        # a multi-byte char cut at the end of the sample is still UTF-8
        if exc.start >= len(sample) - 3:
            try:
                sample[: exc.start].decode("utf-8")
                return "utf-8", False
            except UnicodeDecodeError:
                pass
    from charset_normalizer import from_bytes

    best = from_bytes(sample).best()
    if best is None:
        return "latin-1", False
    enc = best.encoding.replace("_", "-")
    return enc, False


def _sniff_dialect(text: str, override: str | None) -> CsvDialect:
    if override:
        return CsvDialect(delimiter=override)
    lines = [ln for ln in text.splitlines() if ln.strip()][:50]
    sample = "\n".join(lines)
    try:
        d = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim, quote = d.delimiter, d.quotechar or '"'
    except csv.Error:
        counts = {c: sum(ln.count(c) for ln in lines) for c in ",;\t|"}
        delim = max(counts, key=lambda c: counts[c]) if any(counts.values()) else ","
        quote = '"'
    return CsvDialect(delimiter=delim, quotechar=quote)


def _read_csv_rows(
    path: Path, dialect: CsvDialect, *, max_rows: int | None = None
) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    bad = 0
    with path.open("r", encoding=dialect.encoding, errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=dialect.delimiter, quotechar=dialect.quotechar)
        try:
            for row in reader:
                rows.append(row)
                if max_rows is not None and len(rows) >= max_rows:
                    break
        except csv.Error:
            bad += 1
    return rows, bad


def _csv_tables(
    path: Path, opts: IngestOptions, *, full: bool
) -> tuple[CsvDialect, DetectedTable, list[str], list[list[Any]], list[str]]:
    with path.open("rb") as fh:
        head = fh.read(256 * 1024)
    encoding, bom = _sniff_encoding(head) if not opts.encoding else (opts.encoding, False)
    try:
        text = head.decode(
            encoding.replace("-sig", "") if encoding == "utf-8-sig" else encoding, errors="replace"
        )
    except LookupError as exc:
        raise IngestError(f"unknown encoding {encoding!r}") from exc
    if encoding == "utf-8-sig":
        text = text.lstrip("﻿")
    delim = opts.delimiter or ("\t" if path.suffix.lower() in (".tsv", ".tab") else None)
    dialect = _sniff_dialect(text, delim)
    dialect.encoding, dialect.has_bom = encoding, bom
    warnings: list[str] = []
    rows, bad = _read_csv_rows(path, dialect, max_rows=None if full else 2_000)
    if bad:
        warnings.append("the file ended inside a quoted field; the last record may be incomplete")
    grid = [[_clean_cell(v) for v in r] for r in rows]
    width = max((len(r) for r in grid), default=0)
    ragged = sum(1 for r in grid if 0 < len(r) != width and any(v is not None for v in r))
    grid = [r + [None] * (width - len(r)) for r in grid]
    if not grid or width == 0:
        raise IngestError(f"no data found in {path.name}")
    # CSV has a single table: the block spans all non-blank rows
    last_filled = max(i for i, r in enumerate(grid) if any(v is not None for v in r))
    first_filled = min(i for i, r in enumerate(grid) if any(v is not None for v in r))
    b = _Block(first_filled, last_filled, 0, width - 1)
    header_override = (opts.header_row - 1) if opts.header_row else None
    header_idx, first, last, skipped, title = _analyse_block(grid, b, header_override=header_override)
    names, data = _table_from_grid(grid, b, header_idx, first, last)
    table = DetectedTable(
        key="data",
        name_suggestion=suggest_table_name(path.stem),
        header_row=header_idx + 1 if header_idx is not None else None,
        first_data_row=first + 1,
        last_data_row=last + 1,
        first_col=1,
        last_col=width,
        skipped_rows=skipped,
        title=title,
    )
    if ragged:
        table.notes.append(
            f"{ragged} rows have a different number of fields than the widest row; missing fields are empty"
        )
    if header_idx is None:
        table.notes.append("no header row detected; columns named column_1, column_2, ...")
    if not full and len(rows) >= 2_000:
        table.notes.append("inspection is based on the first 2,000 rows")
    return dialect, table, names, data, warnings


# ------------------------------------------------------------------------- JSON


def _json_frame(path: Path, fmt: Format) -> tuple[pd.DataFrame, list[str]]:
    notes: list[str] = []
    if fmt == "ndjson":
        records = []
        with path.open("r", encoding="utf-8-sig") as fh:
            for i, line in enumerate(fh, 1):
                if not line.strip():
                    continue
                try:
                    obj = json.loads(line)
                except ValueError as exc:
                    raise IngestError(f"line {i} is not valid JSON: {exc}") from exc
                if not isinstance(obj, dict):
                    raise IngestError(f"line {i} is not a JSON object")
                records.append(obj)
        data: Any = records
    else:
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except ValueError as exc:
            raise IngestError(f"invalid JSON: {exc}") from exc
    if isinstance(data, dict):
        list_keys = [
            k for k, v in data.items() if isinstance(v, list) and v and all(isinstance(x, dict) for x in v)
        ]
        if len(list_keys) == 1:
            notes.append(f"records read from the {list_keys[0]!r} key")
            data = data[list_keys[0]]
        elif (
            data
            and all(isinstance(v, list) for v in data.values())
            and len({len(v) for v in data.values()}) == 1
        ):
            notes.append("columnar JSON object (one list per column)")
            data = [dict(zip(data.keys(), vals, strict=True)) for vals in zip(*data.values(), strict=True)]
        else:
            data = [data]
            notes.append("single JSON object read as one row")
    if not isinstance(data, list):
        raise IngestError("JSON must be an array of objects, an object wrapping one, or NDJSON")
    if data and not all(isinstance(x, dict) for x in data):
        if all(not isinstance(x, dict | list) for x in data):
            data = [{"value": x} for x in data]
        else:
            raise IngestError("JSON array must contain objects")
    df = pd.json_normalize(data, sep="_") if data else pd.DataFrame()
    for c in df.columns:
        if df[c].map(lambda v: isinstance(v, list | dict)).any():
            df[c] = df[c].map(lambda v: json.dumps(v) if isinstance(v, list | dict) else v)
            notes.append(f"nested values in {c!r} stored as JSON text")
    if any("_" in str(c) for c in df.columns) and any(
        isinstance(x, dict) and any(isinstance(v, dict) for v in x.values()) for x in data[:50]
    ):
        notes.append("nested objects flattened with '_' separators")
    return df, notes


# ------------------------------------------------------------------------ public


def inspect_file(
    path: str | Path,
    *,
    max_bytes: int = DEFAULT_MAX_BYTES,
    options: IngestOptions | None = None,
    preview_rows: int = PREVIEW_ROWS,
) -> FileInspection:
    """Inspect a file without ingesting it: detected tables, header rows, types, preview.

    ``options`` apply exactly as in :func:`ingest_file` (sheet, header row, delimiter,
    encoding, column types), so the preview equals what ingest will load. ``preview_rows``
    caps the preview rows per table (1..1000). Workbooks are streamed and only the first
    ``INSPECT_MAX_ROWS`` rows of each sheet are read.
    """
    opts = options or IngestOptions(max_bytes=max_bytes)
    if options is not None and max_bytes != DEFAULT_MAX_BYTES:
        opts = opts.model_copy(update={"max_bytes": min(max_bytes, opts.max_bytes)})
    preview_rows = max(1, min(int(preview_rows), 1000))
    p = _check_file(path, opts.max_bytes)
    fmt = detect_format(p)
    size = p.stat().st_size
    insp = FileInspection(path=str(p), file_name=p.name, format=fmt, size_bytes=size)
    if fmt == "csv":
        dialect, table, names, data, warnings = _csv_tables(p, opts, full=False)
        df, cols, notes = _frame_from_rows(names, data, column_types=opts.column_types)
        _finish_table(table, df, cols, notes, preview_rows=preview_rows)
        insp.encoding, insp.dialect = dialect.encoding, dialect
        insp.tables, insp.warnings = [table], warnings
        if table.notes and any("first 2,000 rows" in n for n in table.notes):
            table.row_count = _count_csv_records(p, dialect) - (table.first_data_row - 1)
    elif fmt == "excel":
        _check_zip(p, opts.max_bytes)
        sheet = opts.sheet
        if sheet is None and opts.header_row is not None and opts.table_key and "!" in opts.table_key:
            sheet = opts.table_key.split("!", 1)[0]
        sheets, tables, warnings = _excel_tables(
            p, sheet_filter=sheet, header_override=opts.header_row, max_rows=INSPECT_MAX_ROWS
        )
        insp.sheets, insp.warnings = sheets, warnings
        for table, names, data in tables:
            # Same typing as ingest (column_types overrides included), so the preview is what
            # ingest will produce.
            df, cols, notes = _frame_from_rows(names, data, column_types=opts.column_types)
            _finish_table(table, df, cols, notes, preview_rows=preview_rows)
            insp.tables.append(table)
        if not insp.tables:
            insp.warnings.append("no tables detected in any sheet")
    elif fmt in ("json", "ndjson"):
        raw, notes = _json_frame(p, fmt)
        df, cols = _infer_frame(raw)
        table = DetectedTable(
            key="data",
            name_suggestion=suggest_table_name(p.stem),
            header_row=None,
            first_data_row=1,
            last_data_row=len(df),
        )
        _finish_table(table, df, cols, notes)
        insp.encoding = "utf-8"
        insp.tables = [table]
    else:
        import pyarrow.parquet as pq

        try:
            pf = pq.ParquetFile(p)
        except Exception as exc:
            raise IngestError(f"cannot read parquet file: {exc}") from exc
        head = next(pf.iter_batches(batch_size=PREVIEW_ROWS), None)
        prev = head.to_pandas() if head is not None else pd.DataFrame(columns=pf.schema_arrow.names)
        cols = [
            DetectedColumn(
                name=normalize_column_name(f.name, i),
                original_name=f.name,
                inferred_type=_arrow_kind(f.type),
                sample_values=[normalize_value(v) for v in prev[f.name].dropna().head(5).tolist()]
                if f.name in prev
                else [],
            )
            for i, f in enumerate(pf.schema_arrow)
        ]
        table = DetectedTable(
            key="data",
            name_suggestion=suggest_table_name(p.stem),
            first_data_row=1,
            last_data_row=pf.metadata.num_rows,
            columns=cols,
            row_count=pf.metadata.num_rows,
            preview_rows=[
                [normalize_value(v) for v in r]
                for r in prev.astype(object).itertuples(index=False, name=None)
            ],
            last_col=len(cols),
        )
        insp.tables = [table]
    return insp


def _count_csv_records(path: Path, dialect: CsvDialect) -> int:
    n = 0
    with path.open("r", encoding=dialect.encoding, errors="replace", newline="") as fh:
        for row in csv.reader(fh, delimiter=dialect.delimiter, quotechar=dialect.quotechar):
            if any(v.strip() for v in row):
                n += 1
    return n


def _arrow_kind(t: Any) -> Literal["integer", "float", "boolean", "date", "timestamp", "string"]:
    import pyarrow as pa

    if pa.types.is_integer(t):
        return "integer"
    if pa.types.is_floating(t) or pa.types.is_decimal(t):
        return "float"
    if pa.types.is_boolean(t):
        return "boolean"
    if pa.types.is_date(t):
        return "date"
    if pa.types.is_timestamp(t):
        return "timestamp"
    return "string"


def _infer_frame(
    raw: pd.DataFrame, column_types: Mapping[str, str] | None = None
) -> tuple[pd.DataFrame, list[DetectedColumn]]:
    names = _dedupe([normalize_column_name(c, i) for i, c in enumerate(raw.columns)])
    cols: dict[str, pd.Series] = {}
    detected = []
    for name, orig in zip(names, raw.columns, strict=True):
        series, info = infer_column(raw[orig].tolist(), name, str(orig), force=(column_types or {}).get(name))
        cols[name] = series
        detected.append(info)
    return pd.DataFrame(cols), detected


def _finish_table(
    table: DetectedTable,
    df: pd.DataFrame,
    cols: list[DetectedColumn],
    notes: list[str],
    *,
    preview_rows: int = PREVIEW_ROWS,
) -> None:
    table.columns = cols
    table.row_count = len(df)
    table.notes.extend(notes)
    table.last_col = max(table.last_col, table.first_col + len(cols) - 1)
    prev = df.head(preview_rows).astype(object)
    table.preview_rows = [
        [normalize_value(None if _is_missing(v) else v) for v in r]
        for r in prev.itertuples(index=False, name=None)
    ]


def _is_missing(v: Any) -> bool:
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def load_table_frame(
    path: str | Path, options: IngestOptions | None = None
) -> tuple[pd.DataFrame, DetectedTable, FileInspection]:
    """Read the selected table of a file into a typed DataFrame (full data, not a preview)."""
    opts = options or IngestOptions()
    p = _check_file(path, opts.max_bytes)
    fmt = detect_format(p)
    insp = FileInspection(path=str(p), file_name=p.name, format=fmt, size_bytes=p.stat().st_size)
    if fmt == "csv":
        dialect, table, names, data, warnings = _csv_tables(p, opts, full=True)
        df, cols, notes = _frame_from_rows(names, data, column_types=opts.column_types)
        _finish_table(table, df, cols, notes)
        insp.encoding, insp.dialect, insp.tables, insp.warnings = dialect.encoding, dialect, [table], warnings
    elif fmt == "excel":
        _check_zip(p, opts.max_bytes)
        sheet = opts.sheet
        key = opts.table_key
        if key and "!" in key and sheet is None:
            sheet = key.split("!", 1)[0]
        sheets, tables, warnings = _excel_tables(p, sheet_filter=sheet, header_override=opts.header_row)
        insp.sheets, insp.warnings = sheets, warnings
        if not tables:
            raise IngestError("no table detected" + (f" in sheet {sheet!r}" if sheet else ""))
        chosen = None
        if key and opts.header_row is None:
            for t in tables:
                if t[0].key == key:
                    chosen = t
            if chosen is None:
                raise IngestError(f"no detected table {key!r}; available: {[t[0].key for t in tables]}")
        else:
            chosen = tables[0]
        table, names, data = chosen
        df, cols, notes = _frame_from_rows(names, data, column_types=opts.column_types)
        _finish_table(table, df, cols, notes)
        insp.tables = [table]
    elif fmt in ("json", "ndjson"):
        raw, notes = _json_frame(p, fmt)
        df, cols = _infer_frame(raw, opts.column_types)
        table = DetectedTable(
            key="data", name_suggestion=suggest_table_name(p.stem), first_data_row=1, last_data_row=len(df)
        )
        _finish_table(table, df, cols, notes)
        insp.tables = [table]
    else:
        import pyarrow.parquet as pq

        try:
            arrow = pq.read_table(p)
        except Exception as exc:
            raise IngestError(f"cannot read parquet file: {exc}") from exc
        names = _dedupe([normalize_column_name(n, i) for i, n in enumerate(arrow.column_names)])
        arrow = arrow.rename_columns(names)
        df = arrow.to_pandas()
        cols = [
            DetectedColumn(name=n, original_name=o, inferred_type=_arrow_kind(arrow.schema.field(n).type))
            for n, o in zip(names, pq.read_schema(p).names, strict=True)
        ]
        table = DetectedTable(
            key="data", name_suggestion=suggest_table_name(p.stem), first_data_row=1, last_data_row=len(df)
        )
        _finish_table(table, df, cols, [])
        insp.tables = [table]
    if opts.column_names:
        unknown = set(opts.column_names) - set(df.columns)
        if unknown:
            raise IngestError(f"cannot rename unknown columns {sorted(unknown)}")
        new_names = [normalize_column_name(opts.column_names.get(c, c), i) for i, c in enumerate(df.columns)]
        if len(set(new_names)) != len(new_names):
            raise IngestError("column renames produce duplicate names")
        df.columns = new_names
        for c, n in zip(table.columns, new_names, strict=False):
            c.name = n
    return df, table, insp


def ingest_file(store: WorkspaceStore, path: str | Path, options: IngestOptions | None = None) -> TableInfo:
    """Ingest one table of ``path`` into ``store`` and return the created table.

    The source file is only read. Use :func:`inspect_file` first to choose a
    ``table_key`` (Excel), header row and column type overrides.
    """
    opts = options or IngestOptions()
    df, table, _ = load_table_frame(path, opts)
    if df.shape[1] == 0:
        raise IngestError("the selected table has no columns")
    name = opts.table_name or table.name_suggestion
    try:
        validate_table_name(name)
    except StoreError as exc:
        raise IngestError(str(exc)) from exc
    for col in df.columns:
        if df[col].dtype == object:
            df[col] = df[col].map(lambda v: None if _is_missing(v) else v)
    arrow = _frame_to_arrow(df, table.columns)
    try:
        return store.write_table(name, arrow, if_exists=opts.if_exists)
    except StoreError as exc:
        raise IngestError(str(exc)) from exc


def _frame_to_arrow(df: pd.DataFrame, cols: list[DetectedColumn]) -> Any:
    import pyarrow as pa

    types = {c.name: c.inferred_type for c in cols}
    arrays = []
    fields = []
    for name in df.columns:
        s = df[name]
        typ = types.get(name, "string")
        if typ == "date":
            arr = pa.array([v if not _is_missing(v) else None for v in s], type=pa.date32())
        elif typ == "timestamp":
            arr = pa.array(pd.to_datetime(s), type=pa.timestamp("us"))
        elif typ == "integer":
            arr = pa.array(s, type=pa.int64(), from_pandas=True)
        elif typ == "float":
            arr = pa.array(s, type=pa.float64(), from_pandas=True)
        elif typ == "boolean":
            arr = pa.array(s, type=pa.bool_(), from_pandas=True)
        else:
            if s.dtype == object:
                arr = pa.array([None if _is_missing(v) else str(v) for v in s], type=pa.string())
            else:
                arr = pa.Array.from_pandas(s)
        arrays.append(arr)
        fields.append(pa.field(str(name), arr.type))
    return pa.Table.from_arrays(arrays, schema=pa.schema(fields))
