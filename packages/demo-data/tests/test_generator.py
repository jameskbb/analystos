"""Generator: determinism, scale, realism and the deliberate mess."""

from __future__ import annotations

import time

import openpyxl
from analystos_demo import generate


def test_same_seed_same_content_hash(demo, tmp_path):
    t0 = time.process_time()  # CPU time: stable on busy CI hosts, unlike wall-clock time
    fresh = generate(tmp_path / "again", seed=42)
    elapsed = time.process_time() - t0
    assert fresh.content_hash == demo.manifest.content_hash
    assert {f.path: f.sha256 for f in fresh.files} == {f.path: f.sha256 for f in demo.manifest.files}
    # About 12 s wall-clock on an idle laptop (spec budget ~20 s); CPU time includes DuckDB threads.
    assert elapsed < 90, f"generation used {elapsed:.1f}s of CPU"


def test_scale_and_shape(demo, raw):
    m = demo.manifest
    assert m.start_date == "2024-10-01" and m.end_date == "2026-09-30"
    assert m.order_count > 100_000
    assert 300_000 < m.order_line_count < 600_000
    lo, hi = raw.execute("SELECT min(order_date), max(order_date) FROM orders").fetchone()
    assert str(lo) == "2024-10-01" and str(hi) == "2026-09-30"
    assert raw.execute("SELECT count(DISTINCT customer_id) FROM customers").fetchone()[0] == 1500
    assert raw.execute("SELECT count(*) FROM products").fetchone()[0] == 400
    states = {r[0] for r in raw.execute("SELECT DISTINCT state FROM branches").fetchall()}
    assert states == {"TX", "OK"}
    assert raw.execute("SELECT count(*) FROM branches").fetchone()[0] >= 12
    segs = {r[0] for r in raw.execute("SELECT DISTINCT segment FROM customers").fetchall()}
    assert segs == {"Enterprise", "Contractor", "Retail"}
    tiers = {r[0] for r in raw.execute("SELECT DISTINCT tier FROM products").fetchall()}
    assert tiers == {"Entry", "Standard", "Premium"}
    cats = raw.execute("SELECT count(DISTINCT category) FROM products").fetchone()[0]
    assert cats >= 8
    # Every table promised by the brief is present.
    tables = {f.table for f in m.files}
    for t in (
        "customers",
        "orders",
        "order_lines",
        "products",
        "branches",
        "sales_reps",
        "inventory_snapshots",
        "returns",
        "budgets_branch",
        "budgets_category",
        "targets",
        "leads",
        "opportunities",
        "revenue_forecast",
    ):
        assert t in tables


def test_line_arithmetic_is_consistent(raw):
    bad = raw.execute(
        """
        SELECT count(*) FROM order_lines
        WHERE abs(gross_amount - round(unit_list_price * quantity, 2)) > 0.011
           OR abs(net_amount - round(unit_net_price * quantity, 2)) > 0.011
           OR abs(discount_amount - (gross_amount - net_amount)) > 0.011
           OR abs(cost_amount - round(unit_cost * quantity, 2)) > 0.011
        """
    ).fetchone()[0]
    assert bad == 0


def test_deliberate_mess_is_present(raw, scenarios):
    q = raw.execute
    assert q("SELECT count(*) - count(DISTINCT customer_id) FROM customers").fetchone()[0] == 8
    # Duplicates are inactive accounts with no orders, so a customer join cannot inflate revenue.
    dup_orders = q(
        """
        SELECT count(*) FROM orders WHERE customer_id IN
          (SELECT customer_id FROM customers GROUP BY 1 HAVING count(*) > 1)
        """
    ).fetchone()[0]
    assert dup_orders == 0
    assert q("SELECT count(*) FROM customers WHERE region IS NULL").fetchone()[0] > 10
    assert q("SELECT count(DISTINCT region) FROM customers").fetchone()[0] > 5  # casing variants
    assert (
        q("SELECT count(*) FROM customers WHERE try_cast(account_opened AS DATE) IS NULL").fetchone()[0] > 5
    )
    assert q("SELECT count(*) FROM order_lines WHERE quantity < 0").fetchone()[0] == 48
    assert q("SELECT count(*) FROM returns WHERE order_id <> trim(order_id)").fetchone()[0] > 100
    cancelled = q("SELECT count(*) FROM orders WHERE status = 'cancelled'").fetchone()[0]
    assert cancelled > 1000
    cancelled_rev = q(
        "SELECT sum(net_amount) FROM order_lines JOIN orders USING (order_id) WHERE status = 'cancelled'"
    ).fetchone()[0]
    assert cancelled_rev > 100_000, "cancelled orders must carry real amounts so the metric filter matters"
    injection = q("SELECT count(*) FROM customers WHERE customer_name ILIKE 'Ignore previous instructions%'")
    assert injection.fetchone()[0] == 1
    kinds = {i["issue"] for i in scenarios["data_quality_issues"]}
    assert any("duplicate" in k for k in kinds)


def test_budget_workbook_is_messy_but_readable(demo):
    wb = openpyxl.load_workbook(demo.data_dir / "budgets.xlsx", data_only=True)
    assert wb.sheetnames == ["Branch Budget", "Category Budget"]
    ws = wb["Branch Budget"]
    assert "Summit Supply" in str(ws["A1"].value)
    assert ws["A4"].value is None
    assert [c.value for c in ws[5]][:4] == ["Branch ID", "Branch", "Budget Month", "Revenue Budget"]
    values = [r[0] for r in ws.iter_rows(min_row=6, values_only=True)]
    assert None in values, "blank separator row between fiscal years"
    assert str(values[-1]).startswith("Notes:")
    ws2 = wb["Category Budget"]
    assert [c.value for c in ws2[4]] == ["Category", "Budget Month", "Revenue Budget", "Units Budget"]
    assert demo.manifest.file("budgets_branch").header_row == 5
    assert demo.manifest.file("budgets_category").header_row == 4


def test_cache_is_reused_and_validated(demo, tmp_path):
    from analystos_demo import bootstrap_paths

    again = bootstrap_paths(demo.data_dir, seed=42)
    assert again.manifest.content_hash == demo.manifest.content_hash


def test_stories_are_robust_to_the_seed(tmp_path):
    """The stories come from causal knobs plus calibration, not from a lucky seed."""
    import json

    generate(tmp_path / "seed7", seed=7)
    s = json.loads((tmp_path / "seed7" / "scenarios.json").read_text())["stories"]
    rev = s["revenue_decline_aug_2026"]
    assert abs(rev["headline"]["pct_change"] - (-0.118)) <= 0.003
    assert 0.35 <= rev["dallas"]["share_of_decline"] <= 0.65
    assert rev["major_customer"]["rank_by_decline"] == 1
    assert rev["decomposition"]["orders_pct_change"] < rev["decomposition"]["aov_pct_change"] < 0
    assert rev["decomposition"]["asp_pct_change"] < 0
    m = s["margin_compression_q2_2026"]["headline"]
    assert abs(m["revenue_pct_change"]) < 0.01 and m["gross_margin_pct_change_pp"] < -2
    c = s["conversion_decline_q2_2026"]
    assert c["headline"]["conversion_rate_change_pp"] < -1.5 and abs(c["headline"]["leads_pct_change"]) < 0.05
    assert c["concentration"]["share_of_decline_explained"] > 0.6
    assert s["inventory_buildup_2026"]["slow_moving_skus"]["share_of_increase"] > 0.6
    assert s["forecast_miss_q2_2026"]["concentration"]["share_of_miss"] > 0.6
