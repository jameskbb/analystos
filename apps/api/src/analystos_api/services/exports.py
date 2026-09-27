"""Export renderers: CSV, XLSX (openpyxl), Markdown, HTML (print-ready), PDF and JSON.

PDF is produced with WeasyPrint when it is installed (``pip install analystos-api[pdf]``). Without
it the API answers ``409 pdf_unavailable`` and the HTML export (which carries print CSS) is the
documented fallback: open it and use the browser's Print -> Save as PDF.
"""

from __future__ import annotations

import base64
import csv
import html
import io
import json
from typing import Any

from ..errors import ApiError

PRINT_CSS = """
@page { size: A4; margin: 18mm 16mm; }
* { box-sizing: border-box; }
body { font: 11pt/1.45 -apple-system, "Segoe UI", Helvetica, Arial, sans-serif; color: #18181b; margin: 0 auto;
       max-width: 900px; padding: 24px; }
h1 { font-size: 20pt; margin: 0 0 4px; } h2 { font-size: 14pt; margin: 24px 0 8px; border-bottom: 1px solid #e4e4e7;
padding-bottom: 4px; } h3 { font-size: 12pt; margin: 16px 0 6px; }
.meta { color: #71717a; font-size: 9pt; margin-bottom: 16px; }
table { border-collapse: collapse; width: 100%; margin: 8px 0 16px; font-size: 9.5pt; page-break-inside: auto; }
th, td { border: 1px solid #e4e4e7; padding: 4px 6px; text-align: left; vertical-align: top; }
th { background: #f4f4f5; font-weight: 600; } td.num { text-align: right; font-variant-numeric: tabular-nums; }
tr { page-break-inside: avoid; }
.stmt { border-left: 3px solid #71717a; padding: 6px 10px; margin: 8px 0; background: #fafafa; }
.stmt.supported_explanation { border-left-color: #2563eb; }
.stmt.hypothesis { border-left: 3px dashed #d97706; font-style: italic; }
.badge { display: inline-block; font-size: 8.5pt; border: 1px solid #d4d4d8; border-radius: 3px; padding: 0 5px;
         margin-right: 4px; color: #3f3f46; font-style: normal; }
.chips { font-size: 9pt; color: #52525b; } .chips span { margin-right: 10px; }
.kpis { display: flex; flex-wrap: wrap; gap: 10px; } .kpi { border: 1px solid #e4e4e7; padding: 8px 12px; min-width: 150px; }
.kpi .v { font-size: 15pt; font-weight: 600; } .kpi .d { font-size: 9pt; color: #52525b; }
pre, code { font: 9pt/1.4 ui-monospace, Menlo, Consolas, monospace; background: #f4f4f5; }
pre { padding: 8px; white-space: pre-wrap; word-break: break-word; }
img.chart { max-width: 100%; border: 1px solid #e4e4e7; }
.muted { color: #71717a; } @media print { body { padding: 0; } a { color: inherit; text-decoration: none; } }
"""


class PdfUnavailable(ApiError):
    status_code = 409
    code = "pdf_unavailable"


def e(v: Any) -> str:
    return html.escape("" if v is None else str(v))


def fmt_value(v: Any, kind: str | None = None) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int | float):
        if kind == "currency":
            return f"${v:,.0f}"
        if kind == "percent":
            return f"{v * 100:.1f}%"
        if isinstance(v, float):
            return f"{v:,.4g}" if abs(v) < 1000 else f"{v:,.2f}"
        return f"{v:,}"
    return str(v)


def pct(v: Any) -> str:
    return "" if not isinstance(v, int | float) else f"{v * 100:+.1f}%"


# ----------------------------------------------------------------------------------- tabular


def to_csv(columns: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        w.writerow(["" if v is None else _csv_safe(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")


def _csv_safe(v: Any) -> Any:
    """Neutralise spreadsheet formula injection in text cells (=, +, -, @ prefixes)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_xlsx(
    sheets: list[tuple[str, list[str], list[list[Any]]]], about: list[tuple[str, str]] | None = None
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    first = True
    used: set[str] = set()
    for title, columns, rows in sheets:
        name = "".join(ch for ch in title if ch not in "[]:*?/\\")[:31] or "Sheet"
        base, i = name, 2
        while name in used:
            name = f"{base[:28]}_{i}"
            i += 1
        used.add(name)
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_xlsx_value(v) for v in r])
        for idx, col in enumerate(columns, start=1):
            width = max([len(str(col))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r)])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 2), 60)
        ws.freeze_panes = "A2"
    if about:
        ws = wb.create_sheet("About")
        for k, v in about:
            ws.append([k, v])
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_value(v: Any) -> Any:
    if isinstance(v, dict | list):
        return json.dumps(v, default=str)
    if isinstance(v, str) and v[:1] == "=":
        return "'" + v
    return v


def table_md(columns: list[str], rows: list[list[Any]], limit: int = 200) -> str:
    if not columns:
        return ""
    out = [
        "| " + " | ".join(c.replace("|", "\\|") for c in columns) + " |",
        "|" + "|".join("---" for _ in columns) + "|",
    ]
    for r in rows[:limit]:
        out.append("| " + " | ".join(fmt_value(v).replace("|", "\\|") for v in r) + " |")
    if len(rows) > limit:
        out.append(f"\n*{len(rows) - limit} more rows not shown.*")
    return "\n".join(out)


def table_html(columns: list[str], rows: list[list[Any]], limit: int = 500) -> str:
    head = "".join(f"<th>{e(c)}</th>" for c in columns)
    body = "".join(
        "<tr>"
        + "".join(
            f'<td class="num">{e(fmt_value(v))}</td>'
            if isinstance(v, int | float) and not isinstance(v, bool)
            else f"<td>{e(fmt_value(v))}</td>"
            for v in r
        )
        + "</tr>"
        for r in rows[:limit]
    )
    more = f'<p class="muted">{len(rows) - limit} more rows not shown.</p>' if len(rows) > limit else ""
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>{more}"


def md_to_html(text: str) -> str:
    """Small, safe Markdown subset (headings, bold, italics, lists, code blocks, paragraphs); all text escaped."""
    import re

    lines = (text or "").splitlines()
    out: list[str] = []
    in_list = in_code = False
    para: list[str] = []

    def inline(s: str) -> str:
        s = e(s)
        s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
        s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", s)
        return re.sub(r"`(.+?)`", r"<code>\1</code>", s)

    def flush() -> None:
        nonlocal para
        if para:
            out.append("<p>" + " ".join(inline(p) for p in para) + "</p>")
            para = []

    for line in lines:
        if line.strip().startswith("```"):
            flush()
            out.append("</pre>" if in_code else "<pre>")
            in_code = not in_code
            continue
        if in_code:
            out.append(e(line))
            continue
        stripped = line.strip()
        if stripped.startswith(("- ", "* ")) or re.match(r"^\d+\. ", stripped):
            flush()
            if not in_list:
                out.append("<ul>")
                in_list = True
            item = re.sub(r"^(\d+\.|[-*])\s+", "", stripped)
            out.append(f"<li>{inline(item)}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if stripped.startswith("#"):
            flush()
            level = min(len(stripped) - len(stripped.lstrip("#")) + 1, 4)
            out.append(f"<h{level}>{inline(stripped.lstrip('#').strip())}</h{level}>")
        elif not stripped:
            flush()
        elif stripped.startswith(">"):
            flush()
            out.append(f'<div class="stmt">{inline(stripped.lstrip("> "))}</div>')
        else:
            para.append(stripped)
    flush()
    if in_list:
        out.append("</ul>")
    if in_code:
        out.append("</pre>")
    return "\n".join(out)


def html_document(title: str, body: str, meta: str = "") -> str:
    return (
        f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{e(title)}</title>'
        f'<meta name="viewport" content="width=device-width, initial-scale=1"><style>{PRINT_CSS}</style>'
        f'</head><body><h1>{e(title)}</h1><div class="meta">{meta}</div>{body}</body></html>'
    )


def html_to_pdf(document: str) -> bytes:
    try:
        from weasyprint import HTML
    except Exception as exc:  # noqa: BLE001 - ImportError or missing system libraries
        raise PdfUnavailable(
            "PDF rendering needs WeasyPrint (pip install 'analystos-api[pdf]'); export HTML and use the "
            "browser's Print -> Save as PDF instead",
            extra={"fallback": "html"},
        ) from exc
    return HTML(string=document).write_pdf()


def weasyprint_available() -> bool:
    try:
        import weasyprint  # type: ignore[import-not-found]  # noqa: F401
    except Exception:  # noqa: BLE001
        return False
    return True


# ----------------------------------------------------------------------------------- statements / blocks

STATEMENT_LABEL = {
    "observation": "Observation",
    "supported_explanation": "Supported explanation",
    "hypothesis": "Hypothesis",
}
STRENGTH_LABEL = {
    "strong": "Strong evidence",
    "moderate": "Moderate evidence",
    "weak": "Weak evidence",
    "hypothesis_only": "Hypothesis only",
}


def chips_html(ctx: list[dict[str, Any]]) -> str:
    if not ctx:
        return ""
    return (
        '<div class="chips">'
        + "".join(f"<span><b>{e(c.get('label'))}:</b> {e(c.get('value'))}</span>" for c in ctx)
        + "</div>"
    )


def chips_md(ctx: list[dict[str, Any]]) -> str:
    return " · ".join(f"**{c.get('label')}:** {c.get('value')}" for c in ctx or [])


def statement_html(s: dict[str, Any]) -> str:
    st = s.get("statement_type", "observation")
    badges = (
        f'<span class="badge">{e(STATEMENT_LABEL.get(st, st))}</span>'
        f'<span class="badge">{e(STRENGTH_LABEL.get(str(s.get("evidence_strength")), s.get("evidence_strength")))}</span>'
        + (f'<span class="badge">{e(s.get("status"))}</span>' if s.get("status") else "")
    )
    reasons = "".join(f"<li>{e(r)}</li>" for r in s.get("evidence_reasons") or [])
    return (
        f'<div class="stmt {e(st)}">{badges}<div>{e(s.get("statement"))}</div>'
        f"{chips_html(s.get('filter_context') or [])}"
        + (f'<ul class="muted">{reasons}</ul>' if reasons else "")
        + "</div>"
    )


def statement_md(s: dict[str, Any]) -> str:
    st = s.get("statement_type", "observation")
    head = f"> **{STATEMENT_LABEL.get(st, st)}** ({STRENGTH_LABEL.get(str(s.get('evidence_strength')), '')}"
    head += f", {s.get('status')})" if s.get("status") else ")"
    lines = [head, f"> {s.get('statement')}"]
    if s.get("filter_context"):
        lines.append(f"> {chips_md(s['filter_context'])}")
    lines.extend(f"> - {r}" for r in s.get("evidence_reasons") or [])
    return "\n".join(lines)


def _cols(result: dict[str, Any]) -> list[str]:
    return [c["name"] if isinstance(c, dict) else str(c) for c in result.get("columns", [])]


def included_blocks(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Blocks the analyst did not exclude (``excluded: true`` blocks never reach an export, review R-32)."""
    return [b for b in blocks or [] if not b.get("excluded")]


def blocks_to_html(blocks: list[dict[str, Any]], images: dict[str, bytes] | None = None) -> str:
    images = images or {}
    parts: list[str] = []
    for b in included_blocks(blocks):
        t = b.get("type")
        if t == "heading":
            parts.append(f"<h2>{e(b.get('text'))}</h2>")
        elif t in ("narrative", "methodology"):
            if t == "methodology":
                parts.append("<h2>Methodology</h2>")
            parts.append(md_to_html(b.get("markdown", "")))
        elif t == "summary":
            parts.append(f"<h2>{e(b.get('title') or 'Summary')}</h2>")
            for key, label in (
                ("observations", "Observed facts"),
                ("supported_explanations", "Supported explanations"),
                ("hypotheses", "Unresolved hypotheses"),
            ):
                items = b.get(key) or []
                cls = {
                    "observations": "observation",
                    "supported_explanations": "supported_explanation",
                    "hypotheses": "hypothesis",
                }[key]
                parts.append(f"<h3>{label}</h3>")
                parts.append(
                    "".join(f'<div class="stmt {cls}">{e(i.get("text"))}</div>' for i in items)
                    or '<p class="muted">None confirmed.</p>'
                )
        elif t == "finding":
            parts.append(statement_html(b.get("snapshot") or {}))
        elif t == "kpi":
            parts.append(_kpi_html([b]))
        elif t == "kpi_overview":
            parts.append(f"<h2>{e(b.get('title') or 'KPIs')}</h2>" + _kpi_html(b.get("items") or []))
        elif t == "chart":
            parts.append(f"<h3>{e(b.get('title'))}</h3>")
            img = images.get(b.get("id", "")) or images.get(b.get("artifact_id") or "")
            if img:
                parts.append(
                    f'<img class="chart" alt="{e(b.get("title"))}" '
                    f'src="data:image/png;base64,{base64.b64encode(img).decode()}">'
                )
            else:
                parts.append(
                    '<p class="muted">Chart image not captured; the interactive chart is in AnalystOS.</p>'
                )
            prov = b.get("provenance") or {}
            if prov.get("filter_context"):
                parts.append(chips_html(prov["filter_context"]))
        elif t == "table":
            res = b.get("result") or {}
            parts.append(f"<h3>{e(b.get('title'))}</h3>" + table_html(_cols(res), res.get("rows", [])))
            if b.get("note"):
                parts.append(f'<p class="muted">{e(b["note"])}</p>')
        elif t == "anomalies":
            parts.append(f"<h2>{e(b.get('title') or 'Anomalies')}</h2>")
            items = b.get("items") or []
            if items:
                parts.append(
                    table_html(
                        ["date", "value", "expected", "direction", "score"],
                        [
                            [
                                i.get("date"),
                                i.get("value"),
                                i.get("expected"),
                                i.get("direction"),
                                i.get("score"),
                            ]
                            for i in items
                        ],
                    )
                )
            else:
                parts.append(f'<p class="muted">{e(b.get("markdown") or "None detected.")}</p>')
        elif t == "open_questions":
            parts.append(f"<h2>{e(b.get('title') or 'Open questions')}</h2>")
            items = b.get("items") or []
            parts.append(
                "".join(
                    f'<div class="stmt hypothesis">{e(i.get("statement"))} '
                    f'<span class="badge">{e(i.get("status"))}</span></div>'
                    for i in items
                )
                or f'<p class="muted">{e(b.get("markdown"))}</p>'
            )
        elif t == "sources":
            parts.append(
                "<h2>Sources</h2><ul>"
                + "".join(
                    f"<li><b>{e(i.get('label'))}</b> ({e(i.get('kind'))}{', ' + e(i.get('version')) if i.get('version') else ''})"
                    f"{' - ' + e(i.get('detail')) if i.get('detail') else ''}</li>"
                    for i in b.get("items") or []
                )
                + "</ul>"
            )
        elif t == "image":
            img = images.get(b.get("id", ""))
            if img:
                parts.append(
                    f'<img class="chart" src="data:image/png;base64,{base64.b64encode(img).decode()}">'
                )
    return "\n".join(parts)


def _kpi_html(items: list[dict[str, Any]]) -> str:
    cells = []
    for k in items:
        cells.append(
            f'<div class="kpi"><div class="muted">{e(k.get("label"))}</div>'
            f'<div class="v">{e(fmt_value(k.get("value"), k.get("format")))}</div>'
            f'<div class="d">{e(pct(k.get("pct_change")))} vs {e(k.get("baseline_period") or "prior")}'
            f"{' · YoY ' + e(pct(k.get('yoy_pct_change'))) if k.get('yoy_pct_change') is not None else ''}"
            f"</div></div>"
        )
    return f'<div class="kpis">{"".join(cells)}</div>'


def blocks_to_md(title: str, blocks: list[dict[str, Any]]) -> str:
    out = [f"# {title}", ""]
    for b in included_blocks(blocks):
        t = b.get("type")
        if t == "heading":
            out += [f"## {b.get('text')}", ""]
        elif t == "narrative":
            out += [b.get("markdown", ""), ""]
        elif t == "methodology":
            out += ["## Methodology", "", b.get("markdown", ""), ""]
        elif t == "summary":
            out += [f"## {b.get('title') or 'Summary'}", ""]
            for key, label in (
                ("observations", "Observed facts"),
                ("supported_explanations", "Supported explanations"),
                ("hypotheses", "Unresolved hypotheses"),
            ):
                out.append(f"### {label}")
                items = b.get(key) or []
                out += [f"- {i.get('text')}" for i in items] or ["- None confirmed."]
                out.append("")
        elif t == "finding":
            out += [statement_md(b.get("snapshot") or {}), ""]
        elif t in ("kpi", "kpi_overview"):
            items = [b] if t == "kpi" else b.get("items") or []
            if t == "kpi_overview":
                out += [f"## {b.get('title') or 'KPIs'}", ""]
            out += [
                f"- **{k.get('label')}**: {fmt_value(k.get('value'), k.get('format'))} "
                f"({pct(k.get('pct_change'))} vs {k.get('baseline_period') or 'prior'})"
                for k in items
            ]
            out.append("")
        elif t == "chart":
            out += [f"**Chart:** {b.get('title')}", ""]
            prov = b.get("provenance") or {}
            if prov.get("sql"):
                out += ["```sql", prov["sql"], "```", ""]
        elif t == "table":
            res = b.get("result") or {}
            out += [f"### {b.get('title')}", "", table_md(_cols(res), res.get("rows", [])), ""]
            if b.get("note"):
                out += [f"*{b['note']}*", ""]
        elif t == "anomalies":
            out += [f"## {b.get('title') or 'Anomalies'}", ""]
            items = b.get("items") or []
            out += [
                f"- {i.get('date')}: {fmt_value(i.get('value'))} (expected {fmt_value(i.get('expected'))})"
                for i in items
            ] or [b.get("markdown") or "None detected."]
            out.append("")
        elif t == "open_questions":
            out += [f"## {b.get('title') or 'Open questions'}", ""]
            out += [
                f"- *Hypothesis/open:* {i.get('statement')} ({i.get('status')})" for i in b.get("items") or []
            ] or [b.get("markdown") or ""]
            out.append("")
        elif t == "sources":
            out += ["## Sources", ""]
            out += [
                f"- {i.get('label')} ({i.get('kind')}{', ' + i['version'] if i.get('version') else ''})"
                f"{' - ' + i['detail'] if i.get('detail') else ''}"
                for i in b.get("items") or []
            ]
            out.append("")
    return "\n".join(out).rstrip() + "\n"
